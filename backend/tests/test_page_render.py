"""Page view (عرض الصفحة): DOCX -> sanitized HTML."""

from __future__ import annotations

import io
import re
import struct
import zlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.shared import Pt, RGBColor

from app.docx.render import render_docx, render_text
from app.docx.service import extract_plain_text


def _png() -> bytes:
    def chunk(t: bytes, d: bytes) -> bytes:
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 4 for _ in range(4))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 4, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def _rtl(p) -> None:
    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))


def make_docx(path, with_image=True):
    d = Document()
    p = d.add_paragraph()
    _rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("الوقائع")
    r.bold = True
    r.underline = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r.font.name = "Simplified Arabic"
    p2 = d.add_paragraph()
    _rtl(p2)
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p2.add_run("كتاب ")
    p2.add_run("جميل").italic = True
    ppr = p2._p.get_or_add_pPr()
    ppr.append(
        parse_xml(
            '<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:top w:val="single" w:sz="8" w:space="1" w:color="FF0000"/></w:pBdr>'
        )
    )
    ppr.append(
        parse_xml(
            '<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:val="clear" w:color="auto" w:fill="D9D9D9"/>'
        )
    )
    p3 = d.add_paragraph("<script>alert(1)</script> & \"q\"")
    if with_image:
        p3.add_run().add_picture(io.BytesIO(_png()), width=Pt(40))
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text = "خلية"
    t.cell(0, 1).text = "أخرى"
    tcpr = t.cell(0, 0)._tc.get_or_add_tcPr()
    tcpr.append(
        parse_xml(
            '<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:val="clear" w:color="auto" w:fill="FFFF00"/>'
        )
    )
    sect = d.sections[0]
    pb = parse_xml(
        '<w:pgBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'w:offsetFrom="page"><w:top w:val="single" w:sz="4" w:space="24" w:color="auto"/>'
        '<w:left w:val="single" w:sz="4" w:space="24" w:color="auto"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="24" w:color="auto"/>'
        '<w:right w:val="single" w:sz="4" w:space="24" w:color="auto"/></w:pgBorders>'
    )
    sect._sectPr.insert(2, pb)
    d.save(str(path))
    return path


def test_formatting_alignment_borders_tables(tmp_path):
    src = make_docx(tmp_path / "a.docx")
    res = render_docx(src)
    h = res["html"]
    assert 'class="dp-page"' in h and "dp-pageborder" in h
    m = re.search(r'<p dir="rtl" style="([^"]*)"><span style="([^"]*)" lang=ar>الوقائع', h)
    assert m, h[:800]
    pstyle, rstyle = m.groups()
    assert "text-align:center" in pstyle
    assert "font-weight:700" in rstyle and "text-decoration-line:underline" in rstyle
    assert "font-size:18pt" in rstyle and "color:#c00000" in rstyle
    assert "Simplified Arabic" in rstyle and "Amiri" in rstyle  # Arabic fallback stack
    assert "text-align:justify" in h
    assert "border-top:1pt solid #ff0000" in h and "background-color:#d9d9d9" in h
    assert "font-style:italic" in h
    assert "<table" in h and "background-color:#ffff00" in h and "border-top" in h
    assert res["stats"]["images"] == 1
    assert 'width="' not in h  # sizes only via whitelisted style


def test_escaping_and_no_active_content(tmp_path):
    h = render_docx(make_docx(tmp_path / "a.docx"))["html"]
    assert "<script" not in h.lower() and "<a " not in h and "onerror" not in h
    assert "&lt;script&gt;alert(1)&lt;/script&gt; &amp; &quot;q&quot;" in h
    assert "javascript:" not in h and "url(" not in h


def test_image_is_data_uri(tmp_path):
    h = render_docx(make_docx(tmp_path / "a.docx"))["html"]
    m = re.search(r'<img src="(data:image/png;base64,[A-Za-z0-9+/=]+)"', h)
    assert m and "http" not in m.group(1)


def test_image_with_bad_magic_is_skipped(tmp_path):
    d = Document()
    d.add_paragraph("x")
    src = make_docx(tmp_path / "a.docx")
    from zipfile import ZipFile

    out = tmp_path / "bad.docx"
    with ZipFile(src) as zi, ZipFile(out, "w") as zo:
        for i in zi.infolist():
            data = zi.read(i)
            if i.filename.startswith("word/media/"):
                data = b"<svg onload=alert(1)></svg>"
            zo.writestr(i, data)
    res = render_docx(out)
    assert res["stats"]["images"] == 0 and res["stats"]["skipped_images"] == 1
    assert "<img" not in res["html"]


def test_page_geometry(tmp_path):
    res = render_docx(make_docx(tmp_path / "a.docx"))
    assert res["page"]["width_pt"] > 400 and "width:" in res["html"].split(">")[0]


def _diac(src):
    text, _ = extract_plain_text(src)
    return "\n".join(
        ln.replace("كتاب", "كِتَابٌ").replace("الوقائع", "الْوَقَائِعُ") for ln in text.split("\n")
    )


def test_diacritized_text_lands_in_same_runs(tmp_path):
    src = make_docx(tmp_path / "a.docx")
    res = render_docx(src, _diac(src))
    h = res["html"]
    assert res["stats"]["diacritics"]["applied"] >= 2
    # marks in the bold/underlined/centered run, formatting intact
    m = re.search(r'text-align:center[^"]*"><span style="([^"]*)" lang=ar>الْوَقَائِعُ</span>', h)
    assert m and "font-weight:700" in m.group(1)
    assert "كِتَابٌ" in h


def test_changed_letters_render_the_editor_text(tmp_path):
    src = make_docx(tmp_path / "a.docx")
    text = _diac(src).replace("الْوَقَائِعُ", "الوقايع")
    res = render_docx(src, text)
    assert res["stats"]["diacritics"]["edited"] >= 1
    assert "الوقايع" in res["html"] and "الوقائع" not in res["html"]


def test_render_text_fallback_escapes():
    res = render_text("سطر <b>x</b>\nثان", "Amiri", 20)
    assert res["source"] == "text" and "&lt;b&gt;" in res["html"] and "<b>" not in res["html"]
    assert res["html"].count("<p ") == 2


def test_endpoint_docx_and_text_documents(client, tmp_path):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    src = make_docx(tmp_path / "a.docx")
    imp = client.post("/api/docx/import", json={"path": str(src), "project_id": pid}).json()
    did = imp["document_id"]
    r = client.post("/api/docx/render", json={"document_id": did, "diacritized": False})
    assert r.status_code == 200, r.text
    assert r.json()["source"] == "docx" and "font-weight:700" in r.json()["html"]
    client.put(f"/api/documents/{did}", json={"diacritized_content": _diac(src)})
    r2 = client.post("/api/docx/render", json={"document_id": did, "diacritized": True})
    assert "كِتَابٌ" in r2.json()["html"]
    plain = client.post(
        "/api/documents", json={"project_id": pid, "title": "t", "raw_content": "نص حر"}
    ).json()["id"]
    r3 = client.post("/api/docx/render", json={"document_id": plain})
    assert r3.json()["source"] == "text" and "نص حر" in r3.json()["html"]
    assert client.post("/api/docx/render", json={"document_id": 9999}).status_code == 404
