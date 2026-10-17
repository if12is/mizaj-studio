from pathlib import Path

import pytest

from app.legal.tafqeet import (
    TafqeetError,
    amount_words,
    canon,
    number_words,
)


def w(amount, cur="EGP", case="nominative", **kw):
    kw.setdefault("adjective", False)
    return amount_words(amount, cur, case, **kw)["words"]


# Hand-verified table (derived from the grammar, independently of the implementation).
EGP_NOM = [
    ("1", "جُنَيْهٌ وَاحِدٌ"),
    ("2", "جُنَيْهَانِ"),
    ("3", "ثَلَاثَةُ جُنَيْهَاتٍ"),
    ("4", "أَرْبَعَةُ جُنَيْهَاتٍ"),
    ("8", "ثَمَانِيَةُ جُنَيْهَاتٍ"),
    ("10", "عَشَرَةُ جُنَيْهَاتٍ"),
    ("11", "أَحَدَ عَشَرَ جُنَيْهًا"),
    ("12", "اثْنَا عَشَرَ جُنَيْهًا"),
    ("13", "ثَلَاثَةَ عَشَرَ جُنَيْهًا"),
    ("19", "تِسْعَةَ عَشَرَ جُنَيْهًا"),
    ("20", "عِشْرُونَ جُنَيْهًا"),
    ("21", "وَاحِدٌ وَعِشْرُونَ جُنَيْهًا"),
    ("22", "اثْنَانِ وَعِشْرُونَ جُنَيْهًا"),
    ("25", "خَمْسَةٌ وَعِشْرُونَ جُنَيْهًا"),
    ("99", "تِسْعَةٌ وَتِسْعُونَ جُنَيْهًا"),
    ("100", "مِائَةُ جُنَيْهٍ"),
    ("101", "مِائَةٌ وَجُنَيْهٌ وَاحِدٌ"),
    ("102", "مِائَةٌ وَجُنَيْهَانِ"),
    ("103", "مِائَةٌ وَثَلَاثَةُ جُنَيْهَاتٍ"),
    ("111", "مِائَةٌ وَأَحَدَ عَشَرَ جُنَيْهًا"),
    ("120", "مِائَةٌ وَعِشْرُونَ جُنَيْهًا"),
    ("200", "مِائَتَا جُنَيْهٍ"),
    ("201", "مِائَتَانِ وَجُنَيْهٌ وَاحِدٌ"),
    ("300", "ثَلَاثُمِائَةِ جُنَيْهٍ"),
    ("500", "خَمْسُمِائَةِ جُنَيْهٍ"),
    ("800", "ثَمَانِمِائَةِ جُنَيْهٍ"),
    ("999", "تِسْعُمِائَةٍ وَتِسْعَةٌ وَتِسْعُونَ جُنَيْهًا"),
    ("1000", "أَلْفُ جُنَيْهٍ"),
    ("1001", "أَلْفٌ وَجُنَيْهٌ وَاحِدٌ"),
    ("1002", "أَلْفٌ وَجُنَيْهَانِ"),
    ("1100", "أَلْفٌ وَمِائَةُ جُنَيْهٍ"),
    ("1500", "أَلْفٌ وَخَمْسُمِائَةِ جُنَيْهٍ"),
    ("2000", "أَلْفَا جُنَيْهٍ"),
    ("2500", "أَلْفَانِ وَخَمْسُمِائَةِ جُنَيْهٍ"),
    ("3000", "ثَلَاثَةُ آلَافِ جُنَيْهٍ"),
    ("10000", "عَشَرَةُ آلَافِ جُنَيْهٍ"),
    ("11000", "أَحَدَ عَشَرَ أَلْفَ جُنَيْهٍ"),
    ("12000", "اثْنَا عَشَرَ أَلْفَ جُنَيْهٍ"),
    ("21000", "وَاحِدٌ وَعِشْرُونَ أَلْفَ جُنَيْهٍ"),
    ("25000", "خَمْسَةٌ وَعِشْرُونَ أَلْفَ جُنَيْهٍ"),
    ("100000", "مِائَةُ أَلْفِ جُنَيْهٍ"),
    ("100500", "مِائَةُ أَلْفٍ وَخَمْسُمِائَةِ جُنَيْهٍ"),
    ("110000", "مِائَةٌ وَعَشَرَةُ آلَافِ جُنَيْهٍ"),
    ("1000000", "مِلْيُونُ جُنَيْهٍ"),
    ("1000001", "مِلْيُونٌ وَجُنَيْهٌ وَاحِدٌ"),
    ("2000000", "مِلْيُونَا جُنَيْهٍ"),
    ("3000000", "ثَلَاثَةُ مَلَايِينَ جُنَيْهٍ"),
    ("1000000000", "مِلْيَارُ جُنَيْهٍ"),
    ("5000000000", "خَمْسَةُ مِلْيَارَاتِ جُنَيْهٍ"),
    (
        "1234567",
        "مِلْيُونٌ وَمِائَتَانِ وَأَرْبَعَةٌ وَثَلَاثُونَ أَلْفًا وَخَمْسُمِائَةٍ وَسَبْعَةٌ وَسِتُّونَ جُنَيْهًا",
    ),
    (
        "1234567.89",
        "مِلْيُونٌ وَمِائَتَانِ وَأَرْبَعَةٌ وَثَلَاثُونَ أَلْفًا وَخَمْسُمِائَةٍ وَسَبْعَةٌ وَسِتُّونَ جُنَيْهًا" " وَتِسْعَةٌ وَثَمَانُونَ قِرْشًا",
    ),
    ("1,500.50", "أَلْفٌ وَخَمْسُمِائَةِ جُنَيْهٍ وَخَمْسُونَ قِرْشًا"),
    ("٣", "ثَلَاثَةُ جُنَيْهَاتٍ"),  # Arabic-Indic digits accepted
]


@pytest.mark.parametrize(("amount", "expected"), EGP_NOM)
def test_egp_nominative_table(amount, expected):
    assert w(amount) == canon(expected)


OTHER_CASES = [
    ("accusative", "1", "جُنَيْهًا وَاحِدًا"),
    ("accusative", "2", "جُنَيْهَيْنِ"),
    ("accusative", "3", "ثَلَاثَةَ جُنَيْهَاتٍ"),
    ("accusative", "12", "اثْنَيْ عَشَرَ جُنَيْهًا"),
    ("accusative", "20", "عِشْرِينَ جُنَيْهًا"),
    ("accusative", "21", "وَاحِدًا وَعِشْرِينَ جُنَيْهًا"),
    ("accusative", "100", "مِائَةَ جُنَيْهٍ"),
    ("accusative", "200", "مِائَتَيْ جُنَيْهٍ"),
    ("accusative", "300", "ثَلَاثَمِائَةِ جُنَيْهٍ"),
    ("accusative", "1000", "أَلْفَ جُنَيْهٍ"),
    ("genitive", "1", "جُنَيْهٍ وَاحِدٍ"),
    ("genitive", "3", "ثَلَاثَةِ جُنَيْهَاتٍ"),
    ("genitive", "100", "مِائَةِ جُنَيْهٍ"),
    ("genitive", "300", "ثَلَاثِمِائَةِ جُنَيْهٍ"),
    ("genitive", "2000", "أَلْفَيْ جُنَيْهٍ"),
    ("genitive", "21", "وَاحِدٍ وَعِشْرِينَ جُنَيْهًا"),
]


@pytest.mark.parametrize(("case", "amount", "expected"), OTHER_CASES)
def test_cases(case, amount, expected):
    assert w(amount, case=case) == canon(expected)


FEMININE_LBP = [
    ("1", "لِيرَةٌ وَاحِدَةٌ"),
    ("2", "لِيرَتَانِ"),
    ("3", "ثَلَاثُ لِيرَاتٍ"),
    ("8", "ثَمَانِي لِيرَاتٍ"),
    ("10", "عَشْرُ لِيرَاتٍ"),
    ("11", "إِحْدَى عَشْرَةَ لِيرَةً"),
    ("12", "اثْنَتَا عَشْرَةَ لِيرَةً"),
    ("13", "ثَلَاثَ عَشْرَةَ لِيرَةً"),
    ("18", "ثَمَانِيَ عَشْرَةَ لِيرَةً"),
    ("21", "وَاحِدَةٌ وَعِشْرُونَ لِيرَةً"),
    ("22", "اثْنَتَانِ وَعِشْرُونَ لِيرَةً"),
    ("23", "ثَلَاثٌ وَعِشْرُونَ لِيرَةً"),
    ("100", "مِائَةُ لِيرَةٍ"),
    ("101", "مِائَةٌ وَلِيرَةٌ وَاحِدَةٌ"),
    ("103", "مِائَةٌ وَثَلَاثُ لِيرَاتٍ"),
    ("1000", "أَلْفُ لِيرَةٍ"),
    ("3000", "ثَلَاثَةُ آلَافِ لِيرَةٍ"),  # thousand is masculine -> masculine numeral
    ("1003", "أَلْفٌ وَثَلَاثُ لِيرَاتٍ"),
]


@pytest.mark.parametrize(("amount", "expected"), FEMININE_LBP)
def test_feminine_currency_agreement(amount, expected):
    assert w(amount, "LBP") == canon(expected)


OTHER_CURRENCIES = [
    ("5.01", "SAR", "خَمْسَةُ رِيَالَاتٍ وَهَلَلَةٌ وَاحِدَةٌ"),  # هللة is feminine
    ("5.02", "SAR", "خَمْسَةُ رِيَالَاتٍ وَهَلَلَتَانِ"),
    ("5.03", "SAR", "خَمْسَةُ رِيَالَاتٍ وَثَلَاثُ هَلَلَاتٍ"),
    ("5.5", "SAR", "خَمْسَةُ رِيَالَاتٍ وَخَمْسُونَ هَلَلَةً"),
    ("100.25", "SAR", "مِائَةُ رِيَالٍ وَخَمْسٌ وَعِشْرُونَ هَلَلَةً"),
    ("10", "AED", "عَشَرَةُ دَرَاهِمَ"),  # diptote plural
    ("2", "AED", "دِرْهَمَانِ"),
    ("1.500", "KWD", "دِينَارٌ وَاحِدٌ وَخَمْسُمِائَةِ فَلْسٍ"),
    ("0.001", "KWD", "فَلْسٌ وَاحِدٌ"),
    ("0.005", "KWD", "خَمْسَةُ فُلُوسٍ"),
    ("3.250", "KWD", "ثَلَاثَةُ دَنَانِيرَ وَمِائَتَانِ وَخَمْسُونَ فَلْسًا"),
    ("1.234", "JOD", "دِينَارٌ وَاحِدٌ وَثَلَاثَةٌ وَعِشْرُونَ قِرْشًا وَأَرْبَعَةُ فُلُوسٍ"),
    ("1.01", "USD", "دُولَارٌ وَاحِدٌ وَسِنْتٌ وَاحِدٌ"),
    ("15.99", "USD", "خَمْسَةَ عَشَرَ دُولَارًا وَتِسْعَةٌ وَتِسْعُونَ سِنْتًا"),
    ("5", "EUR", "خَمْسَةُ يُورُو"),
    ("1", "EUR", "يُورُو وَاحِدٌ"),
    ("2.50", "OMR", "رِيَالَانِ وَخَمْسُمِائَةِ بَيْسَةٍ"),
]


@pytest.mark.parametrize(("amount", "cur", "expected"), OTHER_CURRENCIES)
def test_currencies(amount, cur, expected):
    assert w(amount, cur) == canon(expected)


def test_adjective_agreement_and_spelling_and_formula():
    assert w("1", adjective=True) == canon("جُنَيْهٌ مِصْرِيٌّ وَاحِدٌ")
    assert w("2", adjective=True) == canon("جُنَيْهَانِ مِصْرِيَّانِ")
    assert w("3", adjective=True) == canon("ثَلَاثَةُ جُنَيْهَاتٍ مِصْرِيَّةٍ")
    assert w("11", adjective=True) == canon("أَحَدَ عَشَرَ جُنَيْهًا مِصْرِيًّا")
    assert w("100", adjective=True) == canon("مِائَةُ جُنَيْهٍ مِصْرِيٍّ")
    assert w("2", "LBP", adjective=True) == canon("لِيرَتَانِ لُبْنَانِيَّتَانِ")
    assert w("100", spelling="mi'a") == canon("مِئَةُ جُنَيْهٍ")
    assert w("200", spelling="mi'a") == canon("مِئَتَا جُنَيْهٍ")
    assert w("300", spelling="mi'a") == canon("ثَلَاثُمِئَةِ جُنَيْهٍ")
    r = amount_words("3", "EGP", legal=True, adjective=False)
    assert r["formula"] == "فَقَطْ ثَلَاثَةُ جُنَيْهَاتٍ لَا غَيْرُ"
    assert r["words"] == "ثَلَاثَةُ جُنَيْهَاتٍ"
    assert amount_words("3", "EGP", adjective=False)["formula"] == r["words"]
    assert w("2.5", minor="fraction", cur="USD") == canon("دُولَارَانِ وَ50/100")


def test_number_words_standalone():
    assert number_words(0) == "صِفْر"
    assert number_words(3) == "ثَلَاثَةٌ"
    assert number_words(1, "f") == "وَاحِدَةٌ"
    assert number_words(200) == "مِائَتَانِ"
    assert number_words(2000) == "أَلْفَانِ"
    assert number_words(21) == "وَاحِدٌ وَعِشْرُونَ"
    assert number_words(2000, case="accusative") == "أَلْفَيْنِ"


@pytest.mark.parametrize(
    ("amount", "cur", "case"),
    [
        ("0", "EGP", None),
        ("0.00", "EGP", None),
        ("-5", "EGP", None),
        ("abc", "EGP", None),
        ("1.234", "EGP", None),  # excess precision is never silently rounded
        ("1.2345", "KWD", None),
        ("5", "XXX", None),
        ("5", "EGP", "dative"),
        ("1" + "0" * 15, "EGP", None),
        ("", "EGP", None),
    ],
)
def test_rejections(amount, cur, case):
    with pytest.raises(TafqeetError):
        amount_words(amount, cur, case)


def test_all_currencies_produce_valid_text():
    from app.engine.normalize import has_arabic_letters
    from app.legal.tafqeet import currencies

    for code in currencies():
        for n in ("1", "2", "7", "11", "100", "1000"):
            out = amount_words(n, code)["words"]
            assert has_arabic_letters(out), code


# ---------- reverse parsing and the amount checker ----------
def test_words_roundtrip_and_pyarabic_reference():
    import random

    from pyarabic import number as pa

    from app.legal.tafqeet import words_to_int

    rng = random.Random(7)
    values = [0, 1, 2, 3, 10, 11, 12, 19, 20, 21, 99, 100, 101, 199, 200, 999, 1000, 1001, 2000]
    values += [rng.randrange(1, 10 ** rng.randrange(1, 13)) for _ in range(300)]
    for n in values:
        for counted in ("m", "f"):
            for case in ("nominative", "accusative", "genitive"):
                text = number_words(n, counted, case)
                assert words_to_int(text.split()) == n, (n, text)
                construct = number_words(n, counted, case, noun_follows=True)
                assert words_to_int(construct.split()) == n, (n, construct)
    # independent reference: pyarabic's spelling of the same value parses to the same value
    for n in [1, 5, 21, 99, 100, 101, 250, 999, 1000, 1500, 12345, 1234567]:
        assert words_to_int(pa.number2text(n).split()) == n, n


def _issues(text):
    from app.legal.amounts import check_amounts

    return check_amounts(text)


def test_check_amounts_matches_and_mismatches():
    ok = "يدفع المستأجر مبلغ 1,500 جنيه (ألف وخمسمائة جنيه) شهرياً."
    assert _issues(ok) == []
    ok2 = "مبلغ وقدره ألف وخمسمائة جنيه مصري (1500) فقط لا غير."
    assert _issues(ok2) == []
    bad = "يدفع المستأجر مبلغ 1,500 جنيه (ألف وخمسمائة وخمسون جنيهاً) شهرياً."
    iss = _issues(bad)
    assert len(iss) == 1 and iss[0]["kind"] == "mismatch"
    assert iss[0]["number_value"] == 1500 and iss[0]["words_value"] == 1550
    assert bad[iss[0]["start"] : iss[0]["end"]].startswith("1,500")
    bad2 = "المبلغ ثلاثة آلاف جنيه (30,000) يسدد فوراً."
    iss2 = _issues(bad2)
    assert len(iss2) == 1 and iss2[0]["words_value"] == 3000 and iss2[0]["number_value"] == 30000
    assert bad2[iss2[0]["start"] : iss2[0]["end"]].startswith("ثلاثة آلاف")


def test_check_amounts_diacritics_decimals_and_noise():
    full = "50.5 جنيه (خَمْسُونَ جُنَيْهًا وَخَمْسُونَ قِرْشًا)"
    assert _issues(full) == []
    assert len(_issues("50.5 جنيه (خمسون جنيهاً وخمسة وخمسون قرشاً)")) == 1
    assert _issues("1001 جنيه (ألف وجنيه واحد)") == []
    assert _issues("2002 جنيه (ألفان وجنيهان)") == []
    assert _issues("البند (5) من العقد 3 أيام") == []  # no words, nothing to compare
    assert _issues("مبلغ 100 (مائة)") == []
    assert len(_issues("مبلغ 100 (مائتان)")) == 1
    # generated words always agree with their figures
    for n in (1, 2, 3, 21, 101, 999, 12000, 1234567):
        t = f"{n} جنيه ({w(str(n))})"
        assert _issues(t) == [], t


# ================================================================ ordinals / clauses
from app.legal.numbering import analyze_clauses, apply_edits, ordinal, parse_ordinal  # noqa: E402

ORD_M = {
    1: "الْأَوَّلُ", 2: "الثَّانِي", 3: "الثَّالِثُ", 4: "الرَّابِعُ", 5: "الْخَامِسُ",
    6: "السَّادِسُ", 7: "السَّابِعُ", 8: "الثَّامِنُ", 9: "التَّاسِعُ", 10: "الْعَاشِرُ",
    11: "الْحَادِيَ عَشَرَ", 12: "الثَّانِيَ عَشَرَ", 13: "الثَّالِثَ عَشَرَ", 19: "التَّاسِعَ عَشَرَ",
    20: "الْعِشْرُونَ", 21: "الْحَادِي وَالْعِشْرُونَ", 22: "الثَّانِي وَالْعِشْرُونَ",
    23: "الثَّالِثُ وَالْعِشْرُونَ", 30: "الثَّلَاثُونَ", 99: "التَّاسِعُ وَالتِّسْعُونَ",
    100: "الْمِائَةُ",
}  # fmt: skip
ORD_F = {
    1: "الْأُولَى", 2: "الثَّانِيَةُ", 3: "الثَّالِثَةُ", 10: "الْعَاشِرَةُ",
    11: "الْحَادِيَةَ عَشْرَةَ", 12: "الثَّانِيَةَ عَشْرَةَ", 13: "الثَّالِثَةَ عَشْرَةَ",
    21: "الْحَادِيَةُ وَالْعِشْرُونَ",
}  # fmt: skip


@pytest.mark.parametrize(("n", "expected"), ORD_M.items())
def test_ordinals_masculine(n, expected):
    assert canon(ordinal(n)) == canon(expected)


@pytest.mark.parametrize(("n", "expected"), ORD_F.items())
def test_ordinals_feminine(n, expected):
    assert canon(ordinal(n, fem=True)) == canon(expected)


def test_ordinal_cases_and_roundtrip():
    assert canon(ordinal(2, case="accusative")) == canon("الثَّانِيَ")
    assert canon(ordinal(5, case="genitive")) == canon("الْخَامِسِ")
    assert canon(ordinal(20, case="genitive")) == canon("الْعِشْرِينَ")
    assert canon(ordinal(21, case="genitive")) == canon("الْحَادِي وَالْعِشْرِينَ")
    for n in range(1, 101):
        for fem in (False, True):
            assert parse_ordinal(ordinal(n, fem)) == n, (n, fem)
            assert parse_ordinal(ordinal(n, fem, diacritized=False)) == n
    assert parse_ordinal("أولاً") == 1 and parse_ordinal("سابعا") == 7
    assert parse_ordinal("كتاب") is None


CONTRACT = """المادة الأولى: التعريفات
المادة الثانية: المدة وفقاً للمادة الرابعة
المادة الرابعة: الالتزامات، وتخضع للمادة (5) أيضاً
المادة الخامسة: الإنهاء وفق المادة الثانية
"""


def test_clauses_gap_renumber_and_references():
    r = analyze_clauses(CONTRACT)
    assert [h["number"] for h in r["headings"]] == [1, 2, 4, 5]
    types = [i["type"] for i in r["issues"]]
    assert "gap" in types and "out_of_order" in types
    fixed = apply_edits(CONTRACT, r["suggestions"])
    assert fixed == CONTRACT.replace("الرابعة: الالتزامات", "الثالثة: الالتزامات").replace(
        "الخامسة: الإنهاء", "الرابعة: الإنهاء"
    ).replace("للمادة الرابعة", "للمادة الثالثة").replace("(5)", "(4)")
    # nothing is applied by the analyzer itself, edits are positional and non-overlapping
    spans = sorted((e["start"], e["end"]) for e in r["suggestions"])
    assert all(a[1] <= b[0] for a, b in zip(spans, spans[1:], strict=False))
    again = analyze_clauses(fixed)
    assert [i["type"] for i in again["issues"] if i["type"] in ("gap", "out_of_order")] == []


def test_clauses_missing_duplicate_circular_external():
    text = """البند الأول: أ
البند الثاني: ب وفقاً للبند الخامس
البند الثاني: ج
المادة 3 من القانون رقم 5 تسري هنا
"""
    r = analyze_clauses(text)
    types = {i["type"] for i in r["issues"]}
    assert "missing_reference" in types and "duplicate_number" in types
    assert not [i for i in r["issues"] if i.get("kind") == "مادة"]  # external law reference
    circ = analyze_clauses("البند الأول: يخضع للبند الثاني\nالبند الثاني: يخضع للبند الأول\n")
    assert "circular_reference" in {i["type"] for i in circ["issues"]}
    clean = analyze_clauses("البند الأول: أ\nالبند الثاني: ب وفق البند الأول\n")
    assert clean["issues"] == [] and clean["suggestions"] == []


def test_clauses_digits_scripts_and_relative():
    text = "المادة (١): أ\nالمادة (٣): ب\nتطبق المادة السابقة\n"
    r = analyze_clauses(text)
    assert r["suggestions"][0]["insert"] == "٢"  # Arabic-Indic digits are preserved
    assert "relative_reference" in {i["type"] for i in r["issues"]}


# ================================================================ dates
from app.legal.dates import date_words  # noqa: E402


def test_dates_both_calendars_and_names():
    r = date_words(2025, 3, 1)
    assert r["hijri"] == {"year": 1446, "month": 9, "day": 1}
    assert r["words"].startswith(canon("يَوْمُ السَّبْتِ الْمُوَافِقُ الْأَوَّلُ مِنْ شَهْرِ مَارِسَ"))
    assert canon("لِسَنَةِ أَلْفَيْنِ وَخَمْسٍ وَعِشْرِينَ مِيلَادِيَّةٍ") in r["words"]
    assert canon("الْمُوَافِقُ الْأَوَّلُ مِنْ شَهْرِ رَمَضَانَ لِسَنَةِ أَلْفٍ وَأَرْبَعِمِائَةٍ وَسِتٍّ وَأَرْبَعِينَ هِجْرِيَّةٍ") in r["words"]
    h = date_words(1447, 9, 15, "hijri")
    assert h["gregorian"] == {"year": 2026, "month": 3, "day": 4}
    assert canon("يَوْمُ الْأَرْبِعَاءِ الْمُوَافِقُ الْخَامِسَ عَشَرَ مِنْ شَهْرِ رَمَضَانَ") in h["words"]
    lv = date_words(2025, 1, 15, gregorian_names="levant", with_day_name=False)
    assert canon("الْخَامِسَ عَشَرَ مِنْ شَهْرِ كَانُونَ الثَّانِي") in lv["words"]
    assert not lv["words"].startswith("يَوْم")
    g = date_words(2025, 3, 1, case="genitive")
    assert g["words"].startswith(canon("يَوْمِ السَّبْتِ الْمُوَافِقِ الْأَوَّلِ"))


def test_dates_adjust_warnings_and_range():
    base = date_words(2025, 3, 1)
    plus = date_words(2025, 3, 1, hijri_adjust=1)
    assert (
        plus["hijri"]["day"] != base["hijri"]["day"]
        or plus["hijri"]["month"] != base["hijri"]["month"]
    )
    assert "+1" in plus["basis"] and base["basis"] == "Umm al-Qura"
    assert (
        date_words(2025, 3, 1, claimed_day_name="الأحد")["warnings"][0]["type"]
        == "day_name_mismatch"
    )
    assert date_words(2025, 3, 1, claimed_day_name="يوم السبت")["warnings"] == []
    with pytest.raises(TafqeetError):
        date_words(1800, 1, 1)  # outside the Umm al-Qura range
    with pytest.raises(TafqeetError):
        date_words(2025, 2, 30)
    with pytest.raises(TafqeetError):
        date_words(2025, 3, 1, "julian")


# ================================================================ defined terms
from app.legal.defined_terms import analyze_terms, canon_term  # noqa: E402

PARTIES = """أبرم هذا العقد بين السيد أحمد محمد (ويشار إليه فيما بعد بـ «الطرف الأول»)، والسيدة فاطمة علي (ويُشار إليها فيما بعد بـ «الطرف الثاني»).
يلتزم الطرف الأول بتسليم العين للطرف الثاني. وإذا أخل الطرف الاول بالتزاماته فللمؤجر الحق في الفسخ. ويقصد بـ «العقار» الشقة المؤجرة. ويحق للطرف 1 الإنهاء.
"""


def test_defined_terms_detection_variants_and_issues():
    r = analyze_terms(PARTIES)
    terms = {t["term"]: t for t in r["terms"]}
    assert set(terms) == {"الطرف الأول", "الطرف الثاني", "العقار"}
    assert terms["الطرف الأول"]["uses"] == 3  # exact, hamza variant, digit variant
    assert terms["الطرف الثاني"]["uses"] == 1
    d = terms["الطرف الأول"]["defined_at"]
    assert PARTIES[d["start"] : d["end"]] == "الطرف الأول"
    kinds = [i["type"] for i in r["issues"]]
    assert kinds.count("variant_form") == 2
    assert {i["found"] for i in r["issues"] if i["type"] == "variant_form"} == {
        "الطرف الاول",
        "للطرف 1",
    }
    assert any(i["type"] == "unused_definition" and i["term"] == "العقار" for i in r["issues"])
    assert any(i["type"] == "undefined_term" and i["term"] == "المؤجر" for i in r["issues"])
    assert any(i["type"] == "name_variants" for i in r["issues"]) is False  # one spelling here


def test_defined_terms_name_variants_duplicates_and_canon():
    t = "السيد أحمد محمد يوقع. ثم السيد احمد محمد يوافق.\nويقصد بـ «العقد» هذا الاتفاق. ويقصد بـ «العقد» أيضاً."
    r = analyze_terms(t)
    assert any(i["type"] == "name_variants" for i in r["issues"])
    assert any(i["type"] == "duplicate_definition" for i in r["issues"])
    assert canon_term("الطرف 1") == canon_term("الطرف الأول") == canon_term("طرف الاول")
    assert analyze_terms("لا توجد تعريفات هنا")["terms"] == []


# ================================================================ glossary + legal level
import sqlite3  # noqa: E402

from app.core.database import connect  # noqa: E402
from app.engine.base import Engine  # noqa: E402
from app.engine.mock_engine import MockEngine  # noqa: E402
from app.engine.normalize import strip_marks  # noqa: E402
from app.engine.pipeline import TashkeelPipeline  # noqa: E402
from app.legal import glossary as gl  # noqa: E402
from app.main import create_app  # noqa: E402


def test_builtin_glossary_is_valid_and_substantial():
    entries = gl.builtin_entries()
    assert len(entries) >= 150
    for e in entries:
        gl.validate_entry(e.term, e.diacritized)  # same letters, no silent letter change
    terms = {e.term for e in entries}
    assert {"العين المؤجرة", "المؤجر", "المستأجر", "الموكل", "يعد", "تعتبر"} <= terms
    by = {e.term: e.diacritized for e in entries}
    assert by["المؤجر"] == "الْمُؤَجِّر" and by["المستأجر"] == "الْمُسْتَأْجِر"
    assert by["الموكل"] == "الْمُوَكِّل" and by["يعد"] == "يُعَدّ"
    with pytest.raises(gl.GlossaryError):
        gl.validate_entry("المؤجر", "الْمُؤَجِّرَة")  # would change a base letter


@pytest.fixture
def conn(data_dir):
    create_app(data_dir)
    c = connect(data_dir / "mizaj.db")
    yield c
    c.close()


def legal_pipeline(engine=None, **kw):
    eng = engine or MockEngine()
    return TashkeelPipeline(
        eng, level="legal", ambiguity=eng.is_ambiguous,
        glossary_loader=lambda c: gl.GlossaryIndex(gl.load_entries(c)), **kw,
    )  # fmt: skip


def test_legal_level_glossary_and_bare_rest(conn):
    out = legal_pipeline().process_chunk(conn, "يلتزم المستأجر بدفع الأجرة").text
    assert out == canon("يلتزم الْمُسْتَأْجِر بدفع الْأُجْرَة")
    # prefixed glossary words: internal marks from the glossary, prefix letters keep the engine's
    out = legal_pipeline().process_chunk(conn, "للمؤجر").text
    assert out.endswith(canon("مُؤَجِّر")) and strip_marks(out) == "للمؤجر"
    out = legal_pipeline().process_chunk(conn, "والمستأجر").text
    assert out.endswith(canon("الْمُسْتَأْجِر"))
    # multi-word terms
    out = legal_pipeline().process_chunk(conn, "تسليم العين المؤجرة").text
    assert out.endswith(canon("الْعَيْن الْمُؤَجَّرَة"))
    # mock ambiguity rule: 3-letter words keep the engine's full marks
    assert legal_pipeline().process_chunk(conn, "عقد").text == "عَقدُ"


class FixedEngine(Engine):
    name, cache_tag = "fixed", "fixed"

    def __init__(self, table):
        self.table = table

    def diacritize(self, words, lookup=None):
        return [self.table.get(w, w) for w in words]

    def is_ambiguous(self, word):
        return word == "علم"


def test_legal_level_keeps_passive_shadda_tanween(conn):
    table = {
        "يعتبر": "يُعْتَبَرُ",
        "ذهب": "ذَهَبَ",
        "مدرس": "مُدَرِّسٌ",
        "كتابا": "كِتَابًا",
        "علم": "عِلْمٌ",
        "بيت": "بَيْتُ",
    }
    eng = FixedEngine(table)
    pipe = TashkeelPipeline(
        eng, level="legal", ambiguity=eng.is_ambiguous,
        passive=lambda bare, dia: dia.startswith("يُ"),
    )  # fmt: skip
    out = pipe.process_chunk(conn, "يعتبر ذهب مدرس كتابا علم بيت").text
    assert out == canon("يُعْتَبَرُ ذهب مدرّسٌ كتابًا عِلْمٌ بيت")  # passive+ambiguous kept
    assert strip_marks(out) == "يعتبر ذهب مدرس كتابا علم بيت"
    with pytest.raises(RuntimeError):
        TashkeelPipeline(eng, level="legal").process_chunk(conn, "علم")


def test_legal_level_user_dictionary_and_lock_win(conn):
    conn.execute(
        "INSERT INTO tashkeel_cache(word_or_phrase, diacritized_form, source, is_user_locked) "
        "VALUES ('المؤجر','الْمُؤْجِرُ','user',1)"
    )
    assert legal_pipeline().process_chunk(conn, "المؤجر").text == "الْمُؤْجِرُ"
    out = legal_pipeline().process_chunk(conn, "المُستأجر").text  # user typed a mark: kept
    assert "مُ" in out and strip_marks(out) == "المستأجر"


def test_user_glossary_overrides_and_disables_builtin(conn):
    from app.core.database import write

    with write(conn):
        gl.save_entry(conn, {"term": "المؤجر", "diacritized": "الْمُؤْجِر"})
        gl.save_entry(conn, {"term": "الضرر", "diacritized": "الضَّرَر", "enabled": False})
        gl.save_entry(conn, {"term": "المنتفع", "diacritized": "الْمُنْتَفِع"})
    entries = {e.term: e for e in gl.load_entries(conn)}
    assert entries["المؤجر"].diacritized == "الْمُؤْجِر" and not entries["المؤجر"].builtin
    assert "الضرر" not in entries and "المنتفع" in entries
    out = legal_pipeline().process_chunk(conn, "المنتفع المؤجر").text
    assert out == canon("الْمُنْتَفِع الْمُؤْجِر")


def test_glossary_avoid_check():
    entries = [gl.Entry("السند لأمر", "السَّنَد لِأَمْر", "السند لأمر", ("السند الأذني",))]
    issues = gl.check_text("هذا السند الأذني مهم", entries)
    assert len(issues) == 1 and issues[0]["preferred"] == "السند لأمر"


def test_real_mishkal_passive_oracle():
    from app.engine.mishkal_engine import MishkalEngine

    eng = MishkalEngine()
    assert eng.is_passive_form("عقد", "عُقِدَ") is True
    assert eng.is_passive_form("عقد", "عَقَدَ") is False
    assert eng.is_passive_form("يعتبر", "يُعْتَبَرُ") is True


# ================================================================ templates
from app.legal.templates import TemplateError, parse, render, validate_template  # noqa: E402

VARS = [
    {"name": "الطرف", "type": "party"},
    {"name": "مبلغ", "type": "money", "currency": "EGP"},
    {"name": "عدد", "type": "number"},
    {"name": "تاريخ", "type": "date"},
    {"name": "نوع", "type": "choice", "choices": ["أ", "ب"], "default": "أ"},
    {"name": "ملاحظة", "type": "text", "required": False},
]


def test_template_rendering_filters_and_conditions():
    body = (
        "الطرف: {{الطرف.name}} ({{الطرف.id_number}})\n"
        "المبلغ: {{amount:مبلغ}} | {{money:مبلغ|legal=true|adjective=false}} | {{figures:مبلغ}}\n"
        "العدد: {{عدد}} = {{words:عدد}}، الترتيبي {{ordinal:عدد|fem=true}}\n"
        '{{#if نوع == "أ"}}نوع أ{{else}}نوع آخر{{/if}}\n'
        "{{#if ملاحظة}}ملاحظة: {{ملاحظة}}{{/if}}تاريخ: {{date:تاريخ|day=false}}"
    )
    vals = {"الطرف": {"name": "أحمد", "id_number": "99"}, "مبلغ": "1500.50", "عدد": 3,
            "تاريخ": "2025-03-01"}  # fmt: skip
    r = render(body, VARS, vals)
    assert r["ok"] and r["errors"] == []
    t = r["text"]
    assert "الطرف: أحمد (99)" in t
    assert "1,500.50 (" in t and t.split("|")[2].strip().startswith("1,500.50")
    assert canon("فَقَطْ أَلْفٌ وَخَمْسُمِائَةِ جُنَيْهٍ وَخَمْسُونَ قِرْشًا لَا غَيْرُ") in t
    assert "العدد: 3 = " + canon("ثَلَاثَةٌ") in t
    assert canon("الثَّالِثَةُ") in t
    assert "نوع أ" in t and "نوع آخر" not in t and "ملاحظة:" not in t
    assert canon("الْمُوَافِقُ الْأَوَّلُ مِنْ شَهْرِ مَارِسَ") in t
    r2 = render(body, VARS, {**vals, "نوع": "ب", "ملاحظة": "x"})
    assert "نوع آخر" in r2["text"] and "ملاحظة: x" in r2["text"]


def test_template_validation_errors_are_reported_not_hidden():
    r = render("{{الطرف.name}} {{amount:مبلغ}}", VARS, {})
    assert not r["ok"] and r["text"] == "[[الطرف]] [[مبلغ]]"
    assert {e["code"] for e in r["errors"]} == {"missing"}
    bad = render("{{amount:مبلغ}} {{words:عدد}} {{date:تاريخ}} {{نوع}}",
                 VARS, {"مبلغ": "abc", "عدد": "x", "تاريخ": "1/2/2025", "نوع": "ج"})  # fmt: skip
    codes = {(e["var"], e["code"]) for e in bad["errors"]}
    assert {("مبلغ", "invalid_amount"), ("عدد", "invalid_value"), ("تاريخ", "invalid_date"),
            ("نوع", "invalid_choice")} <= codes  # fmt: skip
    r3 = render("{{الطرف.name}}", VARS, {"الطرف": {"name": "س"}, "غريب": 1})
    assert [e["code"] for e in r3["errors"]] == ["unknown_variable"]
    # money without currency
    r4 = render("{{amount:m}}", [{"name": "m", "type": "money"}], {"m": "5"})
    assert r4["errors"][0]["code"] == "currency_missing"


@pytest.mark.parametrize(
    "body",
    [
        "{{ __class__ }}",
        "{{ a.__class__ }}",
        "{{ a.name.upper() }}",
        "{{unknown:a}}",
        "{{#if a}}no end",
        "{{/if}}",
        "{{else}}",
        "{{#for x in y}}{{/for}}",
        "{{a|opt}}",
        "{{#if a ==}}x{{/if}}",
        "{{ a b }}",
    ],
)
def test_template_syntax_is_strict_and_safe(body):
    with pytest.raises(TemplateError):
        validate_template(body, [{"name": "a", "type": "party"}])


def test_template_is_not_executable_and_values_are_not_re_evaluated():
    vars_ = [{"name": "t", "type": "text"}]
    r = render("{% set x = 1 %} ${7*7} {{t}} {t}", vars_, {"t": "{{ __import__('os') }}"})
    assert r["ok"]
    assert r["text"] == "{% set x = 1 %} ${7*7} {{ __import__('os') }} {t}"
    with pytest.raises(TemplateError):
        parse("x" * 300_000)
    with pytest.raises(TemplateError):
        render("{{a}}", [{"name": "a", "type": "evil"}], {})
    with pytest.raises(TemplateError):
        render("{{b}}", [{"name": "a", "type": "text"}], {})  # undefined variable


# ================================================================ HTTP API
def test_api_tafqeet_and_currencies(client):
    r = client.post("/api/legal/tafqeet", json={"amount": "1500.50", "currency": "egp"})
    assert r.status_code == 200
    j = r.json()
    assert set(j) == {"words", "formula", "currency", "amount"} and j["currency"] == "EGP"
    assert j["words"] == canon("أَلْفٌ وَخَمْسُمِائَةِ جُنَيْهٍ مِصْرِيٍّ وَخَمْسُونَ قِرْشًا")
    r = client.post("/api/legal/tafqeet", json={"amount": 3, "currency": "SAR", "style": "legal",
                                                 "adjective": False, "case": "genitive"})  # fmt: skip
    assert r.json()["formula"] == canon("فَقَطْ ثَلَاثَةِ رِيَالَاتٍ لَا غَيْرُ")
    assert client.post("/api/legal/tafqeet", json={"amount": 21, "gender": "f"}).json()[
        "words"
    ] == canon("وَاحِدَةٌ وَعِشْرُونَ")
    for bad in ({"amount": "0", "currency": "EGP"}, {"amount": "-1", "currency": "EGP"},
                {"amount": "1.234", "currency": "EGP"}, {"amount": "1", "currency": "XXX"},
                {"amount": "1.5"}, {"amount": "x", "currency": "EGP"}):  # fmt: skip
        r = client.post("/api/legal/tafqeet", json=bad)
        assert r.status_code == 422 and r.json()["error_code"] == "ERR_VALIDATION", bad
    cur = client.get("/api/legal/currencies").json()
    egp = [c for c in cur if c["code"] == "EGP"][0]
    assert egp["subunits"][0]["digits"] == 2 and {c["code"] for c in cur} >= {"SAR", "KWD", "JOD"}


def test_api_check_amounts_dates_terms_clauses(client):
    r = client.post("/api/legal/check-amounts",
                    json={"text": "مبلغ 1,500 جنيه (ألف وخمسمائة وخمسون جنيهاً)"})  # fmt: skip
    assert r.json()["count"] == 1 and r.json()["issues"][0]["kind"] == "mismatch"
    r = client.post("/api/legal/date-words", json={"date": "2025-03-01"})
    assert r.status_code == 200 and r.json()["hijri"]["month"] == 9
    r = client.post("/api/legal/date-words", json={"date": "1447-09-15", "calendar": "hijri",
                                                   "claimed_day_name": "الأحد"})  # fmt: skip
    assert r.json()["warnings"] and r.json()["gregorian"]["day"] == 4
    assert client.post("/api/legal/date-words", json={"date": "1800-01-01"}).status_code == 422
    assert client.post("/api/legal/date-words", json={"date": "oops"}).status_code == 422
    r = client.post("/api/legal/defined-terms", json={"text": PARTIES})
    assert {t["term"] for t in r.json()["terms"]} >= {"الطرف الأول", "الطرف الثاني"}
    r = client.post("/api/legal/clauses", json={"text": CONTRACT})
    assert r.json()["suggestions"] and set(r.json()) == {
        "headings",
        "references",
        "issues",
        "suggestions",
    }


def test_api_glossary_crud_import_export_check(client):
    g = client.get("/api/legal/glossary").json()
    assert len(g) >= 150 and all(e["builtin"] for e in g)
    assert [e["term"] for e in client.get("/api/legal/glossary", params={"q": "مؤجر"}).json()]
    r = client.post("/api/legal/glossary", json={"term": "المنتفع", "diacritized": "الْمُنْتَفِع"})
    assert r.status_code == 201
    eid = r.json()["id"]
    assert not r.json()["builtin"]
    assert (
        client.post(
            "/api/legal/glossary", json={"term": "المنتفع", "diacritized": "الْمُنْتَفِع"}
        ).status_code
        == 422
    )
    bad = client.post("/api/legal/glossary", json={"term": "المنتفع2", "diacritized": "الْمُنْتَفِعُونَ"})
    assert bad.status_code == 422 and bad.json()["error_code"] == "ERR_VALIDATION"
    r = client.put(f"/api/legal/glossary/{eid}", json={"term": "المنتفع", "diacritized": "الْمُنْتَفِع",
                                                       "avoid": ["المستفيد"], "note": "n"})  # fmt: skip
    assert r.json()["avoid"] == ["المستفيد"]
    chk = client.post("/api/legal/glossary/check", json={"text": "يلتزم المستفيد بالسداد"}).json()
    assert chk["count"] == 1 and chk["issues"][0]["preferred"] == "المنتفع"
    exp = client.get("/api/legal/glossary/export").json()
    assert exp["kind"] == "glossary" and [e["term"] for e in exp["entries"]] == ["المنتفع"]
    assert client.delete(f"/api/legal/glossary/{eid}").status_code == 204
    assert client.delete(f"/api/legal/glossary/{eid}").status_code == 404
    r = client.post("/api/legal/glossary/import", json={"entries": exp["entries"]})
    assert r.json() == {"imported": 1, "updated": 0}
    r = client.post("/api/legal/glossary/import", json={"entries": exp["entries"]})
    assert r.json() == {"imported": 0, "updated": 1}


def test_api_legal_level_everywhere(client, tmp_path):
    r = client.post("/api/tashkeel", json={"text": "يلتزم المستأجر بدفع الأجرة", "level": "legal"})
    assert r.status_code == 200 and r.json()["text"] == canon("يلتزم الْمُسْتَأْجِر بدفع الْأُجْرَة")
    r = client.put("/api/settings", json={"tashkeel_level": "legal"})
    assert r.status_code == 200 and r.json()["tashkeel_level"] == "legal"
    # the saved default is honoured when the request has no level
    r = client.post("/api/tashkeel", json={"text": "المؤجر"})
    assert r.json()["text"] == canon("الْمُؤَجِّر")
    s = client.post(
        "/api/tashkeel/stream", json={"text": "المؤجر والمستأجر", "level": "legal"}
    ).text
    assert canon("الْمُؤَجِّر") in s and canon("الْمُسْتَأْجِر") in s
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    did = client.post("/api/documents", json={"project_id": pid, "title": "t",
                                              "raw_content": "المستأجر والمؤجر"}).json()["id"]  # fmt: skip
    out = tmp_path / "legal.docx"
    r = client.post(
        "/api/docx/export", json={"document_id": did, "target_path": str(out), "level": "legal"}
    )
    assert r.status_code == 200
    from docx import Document

    assert canon("الْمُسْتَأْجِر") in Document(str(out)).paragraphs[0].text
    assert client.post("/api/tashkeel", json={"text": "x", "level": "nope"}).status_code == 422


def test_api_templates_seed_crud_render_import_export(client):
    lst = client.get("/api/legal/templates").json()
    assert len(lst) == 6 and all(t["is_sample"] for t in lst)
    names = {t["name"] for t in lst}
    assert {
        "عقد إيجار سكني",
        "عقد بيع منقول",
        "عقد عمل",
        "توكيل خاص (صيغة عامة)",
        "إنذار رسمي",
        "خطاب رسمي",
    } == names
    assert all("body" not in t for t in lst)
    lease = [t for t in lst if t["name"] == "عقد إيجار سكني"][0]
    full = client.get(f"/api/legal/templates/{lease['id']}").json()
    assert "نموذج تعليمي" in full["body"] and "مراجعة مختص" in full["body"]  # sample marking
    # render with missing values -> errors, placeholders stay visible
    r = client.post(f"/api/legal/templates/{lease['id']}/render", json={"values": {}})
    assert r.status_code == 200 and not r.json()["ok"] and "[[" in r.json()["text"]
    assert {e["code"] for e in r.json()["errors"]} == {"missing"}
    vals = {
        "المؤجر": {"name": "أحمد محمد", "nationality": "مصري", "id_number": "1", "address": "القاهرة"},
        "المستأجر": {"name": "سالم علي", "nationality": "مصري", "id_number": "2", "address": "الجيزة"},
        "عنوان_العين": "شارع النيل 5", "وصف_العين": "شقة", "الأجرة_الشهرية": "2500",
        "مبلغ_التأمين": "5000", "تاريخ_البدء": "2025-03-01", "المدة_بالأشهر": 12, "يوم_السداد": 5,
        "شرط_جزائي": "نعم", "مدينة_التحرير": "القاهرة", "تاريخ_التحرير": "2025-02-20",
    }  # fmt: skip
    r = client.post(f"/api/legal/templates/{lease['id']}/render", json={"values": vals})
    j = r.json()
    assert j["ok"], j["errors"]
    assert canon("أَلْفَانِ وَخَمْسُمِائَةِ جُنَيْهٍ مِصْرِيٍّ") in j["text"]
    assert canon("الْخَامِسِ") in j["text"] and "البند السادس" in j["text"]
    assert canon("بِتَارِيخِ")[:0] == "" and canon("يَوْمِ الْخَمِيسِ") in j["text"]
    j2 = client.post(f"/api/legal/templates/{lease['id']}/render",
                     json={"values": {**vals, "شرط_جزائي": "لا", "مبلغ_التأمين": ""}}).json()  # fmt: skip
    assert "البند السادس" not in j2["text"] and "مبلغ التأمين" not in j2["text"]
    # CRUD
    new = {
        "name": "قالب",
        "category": "x",
        "body": "مرحباً {{ن}}",
        "vars": [{"name": "ن", "type": "text"}],
    }
    c = client.post("/api/legal/templates", json=new)
    assert c.status_code == 201 and c.json()["version"] == 1
    tid = c.json()["id"]
    u = client.put(f"/api/legal/templates/{tid}", json={**new, "body": "أهلاً {{ن}}"})
    assert u.json()["version"] == 2 and u.json()["body"] == "أهلاً {{ن}}"
    assert (
        client.post(f"/api/legal/templates/{tid}/render", json={"values": {"ن": "علي"}}).json()[
            "text"
        ]
        == "أهلاً علي"
    )
    bad = client.post("/api/legal/templates", json={**new, "body": "{{مجهول}}"})
    assert bad.status_code == 422 and bad.json()["error_code"] == "ERR_VALIDATION"
    assert client.post("/api/legal/templates", json={**new, "body": "{{#if ن}}"}).status_code == 422
    # export / import round trip
    exp = client.get("/api/legal/templates/export").json()
    assert exp["kind"] == "templates" and len(exp["templates"]) == 7
    assert client.delete(f"/api/legal/templates/{tid}").status_code == 204
    assert client.get(f"/api/legal/templates/{tid}").status_code == 404
    imp = client.post("/api/legal/templates/import", json={"templates": exp["templates"][-1:]})
    assert imp.status_code == 201 and imp.json()["imported"] == 1
    # preview of an unsaved template
    pv = client.post("/api/legal/templates/preview",
                     json={"body": "{{ن}}!", "vars": [{"name": "ن", "type": "text"}], "values": {"ن": "ك"}})  # fmt: skip
    assert pv.json() == {"text": "ك!", "ok": True, "errors": []}
    assert (
        client.get("/api/legal/templates", params={"category": "إيجار"}).json()[0]["name"]
        == "عقد إيجار سكني"
    )


def test_all_starter_templates_validate_and_render(client):
    for t in client.get("/api/legal/templates").json():
        full = client.get(f"/api/legal/templates/{t['id']}").json()
        info = validate_template(full["body"], full["vars"])
        assert info["unused_variables"] == [], t["name"]
        r = client.post(f"/api/legal/templates/{t['id']}/render", json={"values": {}}).json()
        assert not r["ok"] and r["text"]  # nothing crashes with empty values


# ================================================================ migration 002
def test_migration_002_upgrade_keeps_data_and_backs_up(tmp_path):
    from app.core.migrations import discover, run_migrations

    data = tmp_path / "d"
    data.mkdir()
    only_v1 = tmp_path / "v1"
    only_v1.mkdir()
    real = [p for v, p in discover(Path(__file__).resolve().parent.parent / "app" / "migrations")]
    (only_v1 / real[0].name).write_text(real[0].read_text(encoding="utf-8"), encoding="utf-8")
    db = data / "mizaj.db"
    assert run_migrations(db, data / "backups", only_v1) == 1
    c = sqlite3.connect(db)
    c.execute("INSERT INTO projects(name) VALUES ('قبل الترقية')")
    c.commit()
    c.close()
    assert run_migrations(db, data / "backups") == 2  # real migrations: 001 -> 002
    backups = list((data / "backups").glob("mizaj-*-v1.db"))
    assert len(backups) == 1
    old = sqlite3.connect(backups[0])
    assert old.execute("PRAGMA user_version").fetchone()[0] == 1
    assert old.execute("SELECT name FROM projects").fetchone()[0] == "قبل الترقية"
    old.close()
    c = sqlite3.connect(db)
    assert c.execute("SELECT name FROM projects").fetchone()[0] == "قبل الترقية"
    tables = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert {"legal_glossary", "legal_templates", "legal_meta"} <= tables
    c.close()


# ---------- group counts ending in 1 or 2 (101 thousand, 2002002 ...) ----------
SPLIT_SCALE = [
    ("101000", "nominative", "مِائَةُ أَلْفٍ وَأَلْفُ جُنَيْهٍ"),
    ("102000", "nominative", "مِائَةُ أَلْفٍ وَأَلْفَا جُنَيْهٍ"),
    ("201000", "nominative", "مِائَتَا أَلْفٍ وَأَلْفُ جُنَيْهٍ"),
    ("1101000", "nominative", "مِلْيُونٌ وَمِائَةُ أَلْفٍ وَأَلْفُ جُنَيْهٍ"),
    ("101000000", "nominative", "مِائَةُ مِلْيُونٍ وَمِلْيُونُ جُنَيْهٍ"),
    ("102000000", "nominative", "مِائَةُ مِلْيُونٍ وَمِلْيُونَا جُنَيْهٍ"),
    ("1001001", "nominative", "مِلْيُونٌ وَأَلْفٌ وَجُنَيْهٌ وَاحِدٌ"),
    ("2002002", "nominative", "مِلْيُونَانِ وَأَلْفَانِ وَجُنَيْهَانِ"),
    ("102000", "accusative", "مِائَةَ أَلْفٍ وَأَلْفَيْ جُنَيْهٍ"),
    ("101000", "genitive", "مِائَةِ أَلْفٍ وَأَلْفِ جُنَيْهٍ"),
    ("101500", "nominative", "مِائَةُ أَلْفٍ وَأَلْفٌ وَخَمْسُمِائَةِ جُنَيْهٍ"),
]


@pytest.mark.parametrize(("amount", "case", "expected"), SPLIT_SCALE)
def test_scale_groups_ending_in_one_or_two(amount, case, expected):
    assert w(amount, case=case) == canon(expected)


def test_plain_integers_match_for_split_scale_groups():
    from app.legal.tafqeet import words_to_int

    assert number_words(101000) == canon("مِائَةُ أَلْفٍ وَأَلْفٌ")
    assert number_words(102000) == canon("مِائَةُ أَلْفٍ وَأَلْفَانِ")
    assert number_words(1001001) == canon("مِلْيُونٌ وَأَلْفٌ وَوَاحِدٌ")
    for n in (101000, 102000, 201000, 1101000, 101000000, 902000000, 1001001, 2002002):
        for case in ("nominative", "accusative", "genitive"):
            assert words_to_int(number_words(n, "m", case).split()) == n
            assert words_to_int(number_words(n, "m", case, noun_follows=True).split()) == n
