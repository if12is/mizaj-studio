"""DOCX export of a Word-imported document carries the editor's stored text (edits, manual
marks, accepted proofreading) run by run instead of re-running the engine."""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from app.docx.extractor import apply_unit_edit, iter_units
from app.docx.stored import plan
from app.engine.normalize import strip_marks

W_T = qn("w:t")
P1 = "ذَهَبَ الوَلَدُ إِلَى المَدْرَسَةِ"
P2 = "هَذَا كِتَابٌ جَمِيلٌ"
P3 = "الفَقْرَةُ الثَّالِثَةُ هُنَا"


def _build(path):
    d = Document()
    p = d.add_paragraph()
    p.add_run("ذهب ").bold = True
    p.add_run("الولد إلى ")
    p.add_run("المدرسة").italic = True
    d.add_paragraph("هذا كتاب جميل")
    d.add_paragraph("الفقرة الثالثة هنا")
    d.save(str(path))
    return path


def _import(client, src):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    r = client.post("/api/docx/import", json={"path": str(src), "project_id": pid})
    return r.json()["document_id"]


def _export(client, did, out):
    r = client.post("/api/docx/export", json={"document_id": did, "target_path": str(out)})
    assert r.status_code == 200, r.text
    return r.json()["stats"], Document(str(out))


def test_export_uses_stored_diacritized_text_and_keeps_formatting(client, tmp_path):
    did = _import(client, _build(tmp_path / "a.docx"))
    stored = "\n".join([P1, P2, P3])
    client.put(f"/api/documents/{did}", json={"diacritized_content": stored})
    stats, d = _export(client, did, tmp_path / "o.docx")
    assert [p.text for p in d.paragraphs] == [P1, P2, P3]  # exactly the user's marks
    runs = d.paragraphs[0].runs
    assert [r.text for r in runs] == ["ذَهَبَ ", "الوَلَدُ إِلَى ", "المَدْرَسَةِ"]
    assert runs[0].bold and runs[2].italic and not runs[1].bold
    assert stats["from_editor"] == 3 and stats["engine_fallback"] == 0
    assert stats["edited_text"] == 0


def test_export_writes_accepted_base_letter_change(client, tmp_path):
    did = _import(client, _build(tmp_path / "a.docx"))
    # proofreading accepted إلى -> على in paragraph 1; paragraphs 2-3 only have user marks
    stored = "ذَهَبَ الوَلَدُ عَلَى المَدْرَسَةِ\nهذا كتابٌ جميل\nالفقرة الثالثة هنا"
    client.put(f"/api/documents/{did}", json={"diacritized_content": stored})
    stats, d = _export(client, did, tmp_path / "o.docx")
    assert [p.text for p in d.paragraphs] == stored.split("\n")
    runs = d.paragraphs[0].runs
    assert runs[0].text == "ذَهَبَ " and runs[0].bold
    assert runs[1].text == "الوَلَدُ عَلَى " and not runs[1].bold
    assert runs[2].text == "المَدْرَسَةِ" and runs[2].italic
    assert stats["edited_text"] == 1 and stats["engine_fallback"] == 0


def test_rewritten_paragraph_carries_editor_text_not_original(client, tmp_path):
    did = _import(client, _build(tmp_path / "a.docx"))
    # paragraph 2 was completely rewritten in the editor: the editor text still wins
    rewritten = "نَصٌّ آخَرُ تَمَامًا مُخْتَلِفٌ"
    client.put(
        f"/api/documents/{did}", json={"diacritized_content": "\n".join([P1, rewritten, P3])}
    )
    stats, d = _export(client, did, tmp_path / "o.docx")
    assert [p.text for p in d.paragraphs] == [P1, rewritten, P3]
    assert stats["engine_fallback"] == 0 and stats["from_editor"] == 3


def test_inserted_editor_line_becomes_new_paragraph_in_order(client, tmp_path):
    did = _import(client, _build(tmp_path / "a.docx"))
    stored = "\n".join([P1, "", "سطر جديد في المحرر", P2, P3])
    client.put(f"/api/documents/{did}", json={"diacritized_content": stored})
    stats, d = _export(client, did, tmp_path / "o.docx")
    assert [p.text for p in d.paragraphs] == [P1, "سطر جديد في المحرر", P2, P3]
    assert stats["added_paragraphs"] == 1 and stats["engine_fallback"] == 0


def test_raw_edit_without_diacritized_is_exported_then_engine_runs(client, tmp_path):
    did = _import(client, _build(tmp_path / "a.docx"))
    raw = client.get(f"/api/documents/{did}").json()["raw_content"].replace("جميل", "رائع")
    client.put(f"/api/documents/{did}", json={"raw_content": raw})
    stats, d = _export(client, did, tmp_path / "o.docx")
    assert strip_marks(d.paragraphs[1].text) == "هذا كتاب رائع"
    assert d.paragraphs[1].text != "هذا كتاب رائع"  # engine still diacritized it
    assert stats["edited_text"] == 1


def test_blank_stored_content_runs_engine_as_before(client, tmp_path):
    did = _import(client, _build(tmp_path / "a.docx"))
    client.put(f"/api/documents/{did}", json={"raw_content": ""})
    stats, d = _export(client, did, tmp_path / "o.docx")
    assert "from_editor" not in stats
    assert strip_marks(d.paragraphs[1].text) == "هذا كتاب جميل"
    assert d.paragraphs[1].text != "هذا كتاب جميل"


def test_apply_unit_edit_unit_level(tmp_path):
    d = Document(str(_build(tmp_path / "a.docx")))
    unit = next(iter_units(d.element.body, "body"))
    assert apply_unit_edit(unit, "ذَهَبَ الولدان إلى المدرسة")
    ts = [t.text for t in d.element.body.iter(W_T)][:3]
    assert ts == ["ذَهَبَ ", "الولدان إلى ", "المدرسة"]


def test_plan_matches_by_marks_free_text(tmp_path):
    d = Document(str(_build(tmp_path / "a.docx")))
    units = list(iter_units(d.element.body, "body"))
    assert plan(units, "x\n" + P2 + "\ny") == ({0: "x", 1: P2, 2: "y"}, {})
    texts, adds = plan(units, P1 + "\n" + P3)  # middle line deleted
    assert texts == {0: P1, 2: P3} and adds == {}
