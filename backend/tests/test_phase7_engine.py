import sys
import types

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx_factory import read_members

from app.api.tashkeel import low_conf_words, u16_offsets
from app.core.database import connect
from app.engine import config, irab_rules
from app.engine.base import Engine
from app.engine.mock_engine import MockEngine
from app.engine.morph_filter import MorphFilterV2
from app.engine.normalize import MARKS, strip_marks
from app.engine.pipeline import TashkeelPipeline, local_spans, pack_raw, unpack_raw
from app.legal.tafqeet import canon
from app.main import create_app


@pytest.fixture
def conn(data_dir):
    create_app(data_dir)
    c = connect(data_dir / "mizaj.db")
    yield c
    c.close()


class ScriptedEngine(Engine):
    """Returns scripted vocalizations and confidences (word -> (form, word_conf, irab_conf))."""

    name, cache_tag = "scripted", "scripted"

    def __init__(self, table):
        self.table = table

    def diacritize(self, words, lookup=None):
        return [self.table[w][0] for w in words]

    def diacritize_scored(self, words, lookup=None):
        return self.diacritize(words), [self.table[w][1:] for w in words]

    def is_ambiguous(self, word):
        return False


# ---------------------------------------------------------------- 7.1.3 confidence
def test_confidence_off_by_default_and_on_request(client):
    r = client.post("/api/tashkeel", json={"text": "علم كتاب جديد"})
    assert "words" not in r.json()
    r = client.post("/api/tashkeel", json={"text": "علم كتاب جديد", "confidence": True})
    j = r.json()
    assert j["confidence_threshold"] == 0.9
    # mock: 3-letter words are 'unsure' (0.4), the rest 0.95
    flagged = [j["text"][w["start"] : w["end"]] for w in j["words"]]
    assert [strip_marks(f) for f in flagged] == ["علم"]
    assert set(j["words"][0]) == {
        "start", "end", "u16_start", "u16_end", "confidence", "irab_confidence",
    }  # fmt: skip
    assert j["words"][0]["confidence"] == 0.4
    # a higher threshold flags more; a threshold of 0 flags nothing
    j = client.post(
        "/api/tashkeel",
        json={"text": "علم كتاب جديد", "confidence": True, "confidence_threshold": 0.99},
    ).json()
    assert len(j["words"]) == 3
    j = client.post(
        "/api/tashkeel",
        json={"text": "علم كتاب", "confidence": True, "confidence_threshold": 0.0},
    ).json()
    assert j["words"] == []


def test_confidence_offsets_utf16_and_multichunk(client):
    text = "😀 علم " + "كلمة طويلة " * 40 + "علم"
    j = client.post("/api/tashkeel", json={"text": text, "confidence": True}).json()
    out = j["text"]
    assert strip_marks(out) == text and len(j["words"]) == 2
    first, last = j["words"]
    assert strip_marks(out[first["start"] : first["end"]]) == "علم"
    assert first["u16_start"] == first["start"] + 1  # the emoji is 2 UTF-16 units, 1 code point
    assert strip_marks(out[last["start"] : last["end"]]) == "علم"  # offsets valid across chunks


def test_u16_helpers():
    assert u16_offsets("a😀b", [0, 1, 2, 3]) == {0: 0, 1: 1, 2: 3, 3: 4}
    ws = [(0, 3, 0.2, 0.5), (4, 7, None, None), (8, 11, 0.95, 0.9)]
    out = low_conf_words("abc def ghi", ws, 0.9, base=10)
    assert out == [
        {"start": 10, "end": 13, "u16_start": 10, "u16_end": 13, "confidence": 0.2,
         "irab_confidence": 0.5}
    ]  # fmt: skip


def test_sse_confidence_events(client):
    r = client.post("/api/tashkeel/stream", json={"text": "علم كتاب جديد.", "confidence": True})
    first = r.text.split("\n\n")[0]
    assert '"words": [{' in first and '"confidence": 0.4' in first
    plain = client.post("/api/tashkeel/stream", json={"text": "علم كتاب"}).text
    assert '"words"' not in plain.split("event: done")[0]


def test_confidence_survives_the_chunk_cache(conn):
    pipe = TashkeelPipeline(MockEngine())
    a = pipe.process_chunk(conn, "علم كتاب")
    b = pipe.process_chunk(conn, "علم كتاب")
    assert not a.cached and b.cached
    assert a.word_conf == b.word_conf and a.word_conf[0][2] == 0.4
    raw, confs = unpack_raw(pack_raw("كَتابُ", [(0.5, 0.25)]), 1)
    assert raw == "كَتابُ" and confs == [(0.5, 0.25)]
    assert unpack_raw("كَتابُ", 1) == ("كَتابُ", [None])  # old cache rows without confidences


def test_calibration_bins():
    from eval.calibration import calibration

    golds = ["كَتَبَ الْوَلَدُ", "ذَهَبَ"]
    preds = ["كَتَبَ الْوَلَدَ", "ذَهَبَ"]
    wconfs = [[(0, 5, 0.99, 0.9), (6, 12, 0.3, 0.3)], [(0, 5, 0.99, 0.9)]]
    cal = calibration(golds, preds, wconfs)
    rows = {r["bin"]: r for r in cal["bins"]}
    assert rows["0.00-0.50"]["word_error_rate"] == 100.0
    assert rows["0.97-1.00"]["word_error_rate"] == 0.0 and rows["0.97-1.00"]["words"] == 2


# ---------------------------------------------------------------- 7.1.5 context-aware minimal
def test_minimal_keeps_unsure_passive_and_listed_words(conn):
    table = {
        "ذهب": ("ذَهَبَ", 0.99, 0.99),
        "الولد": ("الْوَلَدُ", 0.99, 0.99),
        "غامض": ("غَامِضٌ", 0.5, 0.5),  # engine unsure -> keep
        "يعتبر": ("يُعْتَبَرُ", 0.99, 0.99),  # passive -> keep
        "علم": ("عِلْمٌ", 0.99, 0.99),  # on the always-diacritize list -> keep
    }
    pipe = TashkeelPipeline(
        ScriptedEngine(table),
        level="minimal",
        ambiguity=lambda w: False,
        passive=lambda bare, dia: dia.startswith("يُ"),
        conf_threshold=0.9,
        minimal_words=frozenset({"علم"}),
    )
    out = pipe.process_chunk(conn, "ذهب الولد غامض يعتبر علم").text
    assert out == canon("ذهب الولد غَامِضٌ يُعْتَبَرُ عِلْمٌ")
    assert strip_marks(out) == "ذهب الولد غامض يعتبر علم"


def test_minimal_default_rules_file_and_mock_behaviour(client):
    # mock confidences agree with the mock ambiguity rule, so minimal behaves as before
    r = client.post("/api/tashkeel", json={"text": "علم كتاب ذهب المدرسة", "level": "minimal"})
    assert r.json()["text"] == canon("عَلمُ كتاب ذَهبُ المدرسة")


# ---------------------------------------------------------------- 7.1.4 morph filter v2
class FakeAnalyzer:
    def __init__(self, table):
        self.table = table

    def candidates(self, word):
        return tuple((v, "اسم", f) for v, f in self.table.get(word, ()))


def test_morph_v2_only_changes_unsure_words_and_never_adds_marks():
    an = FakeAnalyzer({"مدرسة": [("مَدْرَسَةٌ", 100), ("مُدَرِّسَةٌ", 10)], "كتاب": [("كِتَابٌ", 5)]})
    mf = MorphFilterV2(an, threshold=0.75)
    # unsure + one vowel away from a candidate -> corrected, ending kept
    assert mf.refine_word("مدرسة", "مُدْرَسَةِ", 0.5) == "مَدْرَسَةِ"
    # confident engine -> untouched
    assert mf.refine_word("مدرسة", "مُدْرَسَةِ", 0.95) == "مُدْرَسَةِ"
    # already a valid analysis -> untouched
    assert mf.refine_word("مدرسة", "مَدْرَسَةٌ", 0.5) == "مَدْرَسَةٌ"
    # candidate has no vowel where the engine has one -> nothing added or removed
    assert mf.refine_word("كتاب", "كُتَابٌ", 0.5) == "كِتَابٌ"
    # unknown words and no confidence -> untouched
    assert mf.refine_word("زززز", "زَزَزَزُ", 0.1) == "زَزَزَزُ"
    assert mf.refine_word("مدرسة", "مُدْرَسَةِ", None) == "مُدْرَسَةِ"
    out = mf.refine(["مدرسة"], ["مُدْرَسَةِ"], [(0.5, 0.5)])
    assert strip_marks(out[0]) == "مدرسة" and out == ["مَدْرَسَةِ"]
    assert mf.replaced >= 3


def test_morph_filter_is_part_of_cache_tag():
    from app.engine.morph_filter import MorphFilter

    base = TashkeelPipeline(MockEngine()).cache_tag
    v2 = TashkeelPipeline(MockEngine(), morph=MorphFilterV2(FakeAnalyzer({}))).cache_tag
    v1 = TashkeelPipeline(MockEngine(), morph=MorphFilter(FakeAnalyzer({}))).cache_tag
    assert len({base, v1, v2}) == 3


# ---------------------------------------------------------------- 7.1.6 i'rab rules
def test_irab_set_case_marks_only():
    pairs = [("ا", ""), ("م", "َ"), ("د", "ُ")]
    fixed = irab_rules.set_case(pairs, "genitive")
    assert fixed[-1] == ("د", "ِ") and [b for b, _ in fixed] == ["ا", "م", "د"]
    assert irab_rules.set_case(pairs, "nominative") is None  # already right
    assert irab_rules.set_case([("ا", ""), ("ب", "ً"), ("ا", "")], "genitive") is None
    assert irab_rules.set_case([("ب", "ٌ")], "genitive")[0] == ("ب", "ٍ")
    assert irab_rules.set_case([("ب", "")], "genitive") is None


def test_irab_plan_rules_and_gates():
    words = ["ذهب", "إلى", "المدرسة", "إن", "الطالب", "كان", "الجو", "جميلا"]
    noun = lambda w: w.startswith("ال")  # noqa: E731
    conf = [0.99, 0.99, 0.5, 0.99, 0.5, 0.99, 0.5, 0.5]
    plan = irab_rules.plan(words, conf, noun, 0.75)
    assert plan == {2: "genitive", 4: "accusative", 6: "nominative"}  # جميلا is not definite
    assert irab_rules.plan(words, [0.99] * 8, noun, 0.75) == {}  # engine confident -> no change
    assert irab_rules.plan(words, conf, None, 0.75) == {}  # no analyzer -> never guess
    assert irab_rules.plan(["بالمدرسة"], [0.2], lambda w: True, 0.75) == {0: "genitive"}


def test_irab_rules_in_the_pipeline(conn):
    table = {
        "ذهب": ("ذَهَبَ", 0.99, 0.99),
        "إلى": ("إِلَى", 0.99, 0.99),
        "المدرسة": ("الْمَدْرَسَةُ", 0.95, 0.4),  # wrong ending, engine unsure
        "الولد": ("الْوَلَدُ", 0.95, 0.99),
    }
    noun = lambda w: w.startswith("ال")  # noqa: E731
    on = TashkeelPipeline(
        ScriptedEngine(table), irab_rules_on=True, irab_noun=noun, use_cache=False
    )
    res = on.process_chunk(conn, "ذهب إلى المدرسة")
    assert res.text.endswith(canon("الْمَدْرَسَةِ")) and res.irab_changes == 1
    off = TashkeelPipeline(ScriptedEngine(table), use_cache=False)
    assert off.process_chunk(conn, "ذهب إلى المدرسة").text.endswith(canon("الْمَدْرَسَةُ"))
    # nothing changes when the engine is sure of the ending
    sure = dict(table, المدرسة=("الْمَدْرَسَةُ", 0.95, 0.99))
    res = TashkeelPipeline(
        ScriptedEngine(sure), irab_rules_on=True, irab_noun=noun, use_cache=False
    ).process_chunk(conn, "ذهب إلى المدرسة")
    assert res.irab_changes == 0 and strip_marks(res.text) == "ذهب إلى المدرسة"


# ---------------------------------------------------------------- Quran protection
def fake_quran(monkeypatch, marker):
    """Install a fake app.quran whose detector protects every occurrence of `marker`."""

    def find(text):
        spans, i = [], text.find(marker)
        while i >= 0:
            spans.append((i, i + len(marker)))
            i = text.find(marker, i + 1)
        return spans

    mod = types.ModuleType("app.quran")
    mod.find_protected_spans = find
    monkeypatch.setitem(sys.modules, "app.quran", mod)
    return find


def test_protected_spans_are_byte_identical(conn):
    protect = lambda t: [(t.index("قل هو"), t.index("قل هو") + 5)]  # noqa: E731
    pipe = TashkeelPipeline(MockEngine(), protector=protect)
    out = pipe.run_text(conn, "كتاب قل هو الله احد جميل")
    assert "قل هو" in out  # protected words keep NO marks added
    assert out.startswith(canon("كَتابُ ")) and out.endswith(canon("جَميلُ"))
    # existing marks inside the span are preserved exactly, even in non-canonical order
    src = "قُلْ هُوَ اللَّهُ أَحَدٌ كتاب"
    protect2 = lambda t: [(0, t.index(" كتاب"))]  # noqa: E731
    out2 = TashkeelPipeline(MockEngine(), protector=protect2).run_text(conn, src)
    assert out2.startswith("قُلْ هُوَ اللَّهُ أَحَدٌ ") and out2.endswith(canon("كَتابُ"))
    # without the protector the same words are diacritized
    free = TashkeelPipeline(MockEngine(), keep_existing=False).run_text(conn, "قل هو")
    assert free != "قل هو"


def test_protected_span_across_chunk_boundary(conn):
    filler = "كلمة عادية " * 30  # > 256 chars: forces several chunks
    holy = "إنا أعطيناك الكوثر"
    text = filler + holy + " " + filler
    start = text.index(holy)
    pipe = TashkeelPipeline(MockEngine(), protector=lambda t: [(start, start + len(holy))])
    out = pipe.run_text(conn, text)
    assert holy in out and strip_marks(out) == text
    assert out.count(holy) == 1
    assert local_spans([(5, 20)], 10, 30) == [(0, 10)] and local_spans([(0, 5)], 10, 30) == []


def test_protector_failures_never_break_diacritization(conn):
    def boom(_):
        raise RuntimeError("detector crashed")

    out = TashkeelPipeline(MockEngine(), protector=boom).run_text(conn, "كتاب جديد")
    assert out == canon("كَتابُ جَديدُ")


def test_quran_module_is_picked_up_by_the_api(client, monkeypatch):
    fake_quran(monkeypatch, "الحمد لله")
    r = client.post("/api/tashkeel", json={"text": "الحمد لله كتاب جديد"})
    assert r.json()["text"].startswith("الحمد لله ")
    assert r.json()["text"].endswith(canon("كَتابُ جَديدُ"))
    r = client.post("/api/tashkeel", json={"text": "الحمد لله كتاب", "protect_quran": False})
    assert r.json()["text"] != "الحمد لله كتاب" and "الحمد لله" not in r.json()["text"]
    s = client.post("/api/tashkeel/stream", json={"text": "الحمد لله كتاب"}).text
    assert '"text": "الحمد لله ' in s  # SSE chunks honour it too


def test_quran_protection_in_docx_export(client, tmp_path, monkeypatch):
    fake_quran(monkeypatch, "قل هو الله احد")
    src = tmp_path / "q.docx"
    d = Document()
    p = d.add_paragraph()
    p.add_run("كتاب جديد قل هو ")
    p.add_run("الله احد")  # the protected span crosses a run boundary
    p.add_run(" وكتاب")
    d.save(str(src))
    out = tmp_path / "out.docx"
    r = client.post("/api/docx/export", json={"source_path": str(src), "target_path": str(out)})
    assert r.status_code == 200
    root = parse_xml(read_members(out)["word/document.xml"])
    texts = [t.text for t in root.iter(qn("w:t"))]
    joined = "".join(texts)
    assert "قل هو الله احد" in joined and strip_marks(joined) == "كتاب جديد قل هو الله احد وكتاب"
    assert joined.startswith(canon("كَتابُ جَديدُ ")) and joined.endswith(canon("وكَتابُ")[-4:])
    assert sum(c in MARKS for c in joined) > 0
    # same invariant text per w:t node
    assert [strip_marks(t) for t in texts] == ["كتاب جديد قل هو ", "الله احد", " وكتاب"]


def test_poetry_mode_keeps_one_chunk_per_line(client):
    from app.engine.splitter import ArabicSentenceSplitter

    text = "سطر أول\nسطر ثان قصير\n\nسطر رابع"
    chunks = ArabicSentenceSplitter(line_mode=True).split(text)
    assert [c.text for c in chunks] == ["سطر أول\n", "سطر ثان قصير\n", "\n", "سطر رابع"]
    assert "".join(c.text for c in chunks) == text
    long_line = "كلمة " * 100
    sub = ArabicSentenceSplitter(line_mode=True).split(long_line + "\nب")
    assert len(sub) >= 3 and "".join(c.text for c in sub) == long_line + "\nب"
    r = client.post("/api/tashkeel/stream", json={"text": text, "poetry": True})
    assert r.text.count('"chunk_id"') == 4


# ---------------------------------------------------------------- 7.1.2 INT8 selection
def test_int8_encoder_selection(tmp_path, monkeypatch):
    d = tmp_path / "catt_eo"
    d.mkdir()
    (d / "encoder.onnx").write_bytes(b"f")
    monkeypatch.setattr(config, "CATT_DIR", d)
    monkeypatch.delenv("MIZAJ_MODELS_DIR", raising=False)
    monkeypatch.delenv("MIZAJ_CATT_PRECISION", raising=False)
    assert config.catt_encoder_path().name == "encoder.onnx" and not config.int8_available()
    (d / "encoder_int8.onnx").write_bytes(b"q")
    assert config.int8_available()
    assert config.catt_encoder_path().name == "encoder.onnx"  # auto prefers float when present
    monkeypatch.setenv("MIZAJ_CATT_PRECISION", "float")
    assert config.catt_encoder_path().name == "encoder.onnx"
    monkeypatch.setenv("MIZAJ_CATT_PRECISION", "int8")
    assert config.catt_encoder_path().name == "encoder_int8.onnx"


def test_int8_only_install_is_found(tmp_path, monkeypatch):
    """An INT8-only model dir (no float encoder) is still found and used in auto mode."""
    d = tmp_path / "catt_eo"
    d.mkdir()
    (d / "encoder_int8.onnx").write_bytes(b"q")
    (d / "decoder.onnx").write_bytes(b"d")
    monkeypatch.setattr(config, "CATT_DIR", tmp_path / "nowhere")
    monkeypatch.setenv("MIZAJ_MODELS_DIR", str(tmp_path))
    monkeypatch.delenv("MIZAJ_CATT_PRECISION", raising=False)
    assert config.resolve_model_dir() == ("env", d)
    assert config.catt_encoder_path().name == "encoder_int8.onnx"  # auto picks INT8
    assert config.model_present()
    monkeypatch.setenv("MIZAJ_CATT_PRECISION", "float")  # explicit float: absent -> not present
    assert not config.model_present()
