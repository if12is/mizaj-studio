"""Builds a rich .docx (python-docx + lxml injection) for the Phase 3 tests."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.shared import RGBColor

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = (
    f'xmlns:w="{W}" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)

FOOTNOTE_TEXT = "هذه حاشية سفلية عن الكتاب"
ENDNOTE_TEXT = "هذه تعليق ختامي عن العلم"
TEXTBOX_TEXT = "نص داخل مربع جميل"
SDT_TEXT = "محتوى عنصر التحكم الجديد"
COMMENT_TEXT = "تعليق المراجع على الكتاب"
HEADER_TEXTS = {
    "default": "ترويسة عادية للكتاب",
    "first": "ترويسة الصفحة الأولى",
    "even": "ترويسة الصفحات الزوجية",
}
FOOTER_TEXTS = {
    "default": "تذييل عادي للكتاب",
    "first": "تذييل الصفحة الأولى",
    "even": "تذييل الصفحات الزوجية",
}
CODE_TEXT = "كتاب calculate_total()"
INSTR_TEXT = " PAGE \\* MERGEFORMAT "


def _x(xml: str):
    return parse_xml(xml)


def _p(*runs_xml: str) -> str:
    return f"<w:p {NS}>{''.join(runs_xml)}</w:p>"


def _r(text: str, rpr: str = "") -> str:
    return f'<w:r><w:rPr>{rpr}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'


def _add_part(doc, name: str, content_type: str, reltype: str, xml: str) -> None:
    part = Part(PackURI(name), content_type, xml.encode("utf-8"), doc.part.package)
    doc.part.relate_to(part, reltype)


def build_rich_docx(path: Path) -> Path:
    doc = Document()
    body = doc.element.body
    sect = body.sectPr

    def add(xml: str):
        el = _x(xml)
        sect.addprevious(el)
        return el

    # 1. word split across runs: identical formatting, then differing formatting
    add(_p(_r("الك"), _r("تاب"), _r(" جميل "), _r("مدر", "<w:b/>"), _r("سة")))
    add(_p(_r("ملون", '<w:color w:val="FF0000"/><w:i/>'), _r(" نص عادي هنا")))
    # 2. hyperlink + bookmark
    add(
        _p(
            '<w:bookmarkStart w:id="7" w:name="mark1"/>',
            _r("قبل الرابط "),
            f'<w:hyperlink w:anchor="mark1">{_r("رابط الموقع")}</w:hyperlink>',
            '<w:bookmarkEnd w:id="7"/>',
        )
    )
    # 3. complex field (instrText must never change) + fldSimple
    add(
        _p(
            _r("الصفحة "),
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
            f'<w:r><w:instrText xml:space="preserve">{INSTR_TEXT}</w:instrText></w:r>',
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>',
            _r("العدد"),
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
            f'<w:fldSimple w:instr=" TITLE ">{_r("العنوان")}</w:fldSimple>',
        )
    )
    # 4. tab / br inside runs
    add(
        _p(
            "<w:r><w:t>قبل</w:t><w:tab/><w:t>بعد</w:t><w:br/><w:t>سطر جديد</w:t></w:r>",
        )
    )
    # 5. code-font runs (skipped) and a mixed paragraph
    rf = '<w:rFonts w:ascii="Courier New" w:hAnsi="Courier New"/>'
    rf2 = '<w:rFonts w:cs="CONSOLAS"/>'
    add(_p(_r("اقرأ الكتاب ثم "), _r(CODE_TEXT, rf), _r(" وبعده "), _r("اختبار", rf2)))
    add(_p(_r("اشتريت جهاز iPhone 16 Pro الجديد بمبلغ 1200 USD من https://example.com/a")))
    # 6. footnote + endnote references, comment anchors
    add(
        _p(
            _r("نص مع حاشية"),
            '<w:r><w:footnoteReference w:id="1"/></w:r>',
            _r(" ونص مع تعليق ختامي"),
            '<w:r><w:endnoteReference w:id="1"/></w:r>',
            '<w:commentRangeStart w:id="0"/>',
            _r(" ونص معلق عليه"),
            '<w:commentRangeEnd w:id="0"/>',
            '<w:r><w:commentReference w:id="0"/></w:r>',
        )
    )
    # 7. text box in mc:AlternateContent (Choice + Fallback, both must be diacritized)
    txbx = f"<w:txbxContent><w:p><w:r><w:t>{TEXTBOX_TEXT}</w:t></w:r></w:p></w:txbxContent>"
    add(
        _p(
            _r("قبل المربع "),
            "<w:r><mc:AlternateContent>"
            f'<mc:Choice Requires="wps"><wps:wsp><wps:txbx>{txbx}</wps:txbx></wps:wsp></mc:Choice>'
            f"<mc:Fallback><w:pict><v:shape><v:textbox>{txbx}</v:textbox></v:shape></w:pict>"
            "</mc:Fallback></mc:AlternateContent></w:r>",
            _r(" بعد المربع"),
        )
    )
    # 8. content control (block level sdt)
    add(
        f'<w:sdt {NS}><w:sdtPr><w:alias w:val="c"/></w:sdtPr><w:sdtContent>'
        f"<w:p><w:r><w:t>{SDT_TEXT}</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    )

    # 9. table: merged cells + nested table
    t = doc.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(0, 0).text = "رأس مدمج للجدول"
    t.cell(0, 2).text = "عمود ثالث"
    t.cell(1, 0).text = "خلية عادية"
    nested = t.cell(1, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "جدول متداخل"
    nested.cell(0, 1).text = "خلية داخلية"
    t.cell(1, 2).paragraphs[0].add_run("غامق").bold = True
    r = t.cell(1, 2).paragraphs[0].add_run(" وملون")
    r.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    # python-docx puts the table before sectPr already (add_table), keep order

    # 10. headers / footers (default, first page, even)
    s = doc.sections[0]
    s.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True
    s.header.paragraphs[0].text = HEADER_TEXTS["default"]
    s.first_page_header.paragraphs[0].text = HEADER_TEXTS["first"]
    s.even_page_header.paragraphs[0].text = HEADER_TEXTS["even"]
    s.footer.paragraphs[0].text = FOOTER_TEXTS["default"]
    s.first_page_footer.paragraphs[0].text = FOOTER_TEXTS["first"]
    s.even_page_footer.paragraphs[0].text = FOOTER_TEXTS["even"]

    # 11. footnotes / endnotes / comments parts
    fn = (
        f'<w:footnotes {NS}><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/>'
        '</w:r></w:p></w:footnote><w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r>'
        '<w:continuationSeparator/></w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:r>'
        "<w:footnoteRef/></w:r>"
        f'<w:r><w:t xml:space="preserve"> {FOOTNOTE_TEXT}</w:t></w:r>'
        "</w:p></w:footnote></w:footnotes>"
    )
    en = (
        f'<w:endnotes {NS}><w:endnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/>'
        '</w:r></w:p></w:endnote><w:endnote w:id="1"><w:p><w:r><w:endnoteRef/></w:r>'
        f'<w:r><w:t xml:space="preserve"> {ENDNOTE_TEXT}</w:t></w:r></w:p></w:endnote></w:endnotes>'
    )
    cm = (
        f'<w:comments {NS}><w:comment w:id="0" w:author="م" w:date="2024-01-01T00:00:00Z">'
        f"<w:p><w:r><w:t>{COMMENT_TEXT}</w:t></w:r></w:p></w:comment></w:comments>"
    )
    _add_part(doc, "/word/footnotes.xml", CT.WML_FOOTNOTES, RT.FOOTNOTES, fn)
    _add_part(doc, "/word/endnotes.xml", CT.WML_ENDNOTES, RT.ENDNOTES, en)
    _add_part(doc, "/word/comments.xml", CT.WML_COMMENTS, RT.COMMENTS, cm)

    doc.save(str(path))
    return path


def read_members(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {i.filename: z.read(i) for i in z.infolist()}


def build_tab_marks_docx(path: Path) -> Path:
    """Marks typed right after a w:tab (they used to be dropped on export)."""
    doc = Document()
    sect = doc.element.body.sectPr
    sect.addprevious(
        _x(_p("<w:r><w:t>كتاب</w:t><w:tab/><w:t>ّعلم</w:t><w:tab/><w:t>َقلم</w:t></w:r>"))
    )
    doc.save(str(path))
    return path
