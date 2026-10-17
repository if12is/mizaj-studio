import asyncio
import json
import os
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx_factory import build_rich_docx, build_tab_marks_docx, read_members

from app.engine import config
from app.engine.manager import manager
from app.engine.normalize import MARKS, strip_marks

BACKEND = Path(__file__).resolve().parent.parent
W_T = qn("w:t")


# ---------- 7.7.5: nothing is written under $HOME ----------
def test_no_home_directory_writes(tmp_path):
    home = tmp_path / "home"
    home.mkdir()
    env = {**os.environ, "HOME": str(home), "USERPROFILE": str(home)}
    env.pop("MIZAJ_ENGINE_MODE", None)
    proc = subprocess.Popen(
        [sys.executable, "-m", "app.main", "--port", "0", "--data-dir", str(tmp_path / "data")],
        cwd=BACKEND,
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
    )
    try:
        port = int(proc.stdout.readline().split("=")[1])
        base = f"http://127.0.0.1:{port}"

        def post(path, body):
            req = urllib.request.Request(
                base + path, json.dumps(body).encode("utf-8"), {"Content-Type": "application/json"}
            )
            return json.load(urllib.request.urlopen(req, timeout=120))

        for _ in range(60):  # wait for warm-up/503
            try:
                out = post("/api/tashkeel", {"text": "ذهب الولد إلى المدرسة", "mode": "fast"})
                break
            except urllib.error.HTTPError as exc:
                assert exc.code == 503
                time.sleep(0.5)
        assert out["engine"] == "mishkal" and out["text"] != "ذهب الولد إلى المدرسة"
        with urllib.request.urlopen(
            base + "/api/word-suggestions?word=%D8%B9%D9%84%D9%85", timeout=60
        ) as r:
            assert json.load(r)
    finally:
        proc.terminate()
        proc.wait(timeout=10)
    assert list(home.iterdir()) == []  # nothing created in HOME
    assert (tmp_path / "data" / "cache" / ".thaalabCache").exists()


# ---------- 7.7.4: warm-up and /health engines ----------
def test_health_engines_states_and_warmup(client):
    h = client.get("/health").json()
    assert h["engines"] == {"fast": "loading", "accurate": "loading"}
    asyncio.run(manager.warmup())
    h = client.get("/health").json()
    assert h["engines"] == {"fast": "ready", "accurate": "ready"}
    # first request after warm-up never sees a 503
    assert client.post("/api/tashkeel", json={"text": "كتاب"}).status_code == 200


def test_health_engine_missing_when_model_absent(client, monkeypatch, tmp_path):
    monkeypatch.delenv("MIZAJ_ENGINE_MODE", raising=False)
    monkeypatch.setattr(config, "CATT_DIR", tmp_path / "none")
    assert client.get("/health").json()["engines"]["accurate"] == "missing"
    assert client.get("/health").json()["engines"]["fast"] == "loading"


def test_handshake_not_delayed_and_engines_become_ready(tmp_path):
    env = {**os.environ, "MIZAJ_ENGINE_MODE": "mock"}
    proc = subprocess.Popen(
        [sys.executable, "-m", "app.main", "--port", "0", "--data-dir", str(tmp_path / "d")],
        cwd=BACKEND,
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
    )
    try:
        port = int(proc.stdout.readline().split("=")[1])
        assert proc.stdout.readline().strip() == "MIZAJ_STATUS=READY"
        engines = {}
        for _ in range(40):
            with urllib.request.urlopen(f"http://127.0.0.1:{port}/health", timeout=5) as r:
                engines = json.load(r)["engines"]
            if set(engines.values()) == {"ready"}:
                break
            time.sleep(0.25)
        assert engines == {"fast": "ready", "accurate": "ready"}
    finally:
        proc.terminate()
        proc.wait(timeout=10)


# ---------- 7.3.1: DOCX for documents without a source file ----------
def _doc(client, **kw):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    r = client.post("/api/documents", json={"project_id": pid, "title": "عنوان", **kw})
    return r.json()["id"]


def test_generate_docx_from_diacritized_content(client, tmp_path):
    did = _doc(client, raw_content="سطر أول\n\nسطر ثالث")
    text = "سَطْرٌ أَوَّلُ\n\nسَطْرٌ ثَالِثُ\nEnglish 123"
    client.put(f"/api/documents/{did}", json={"diacritized_content": text})
    client.put(
        "/api/settings", json={"editor_font": "Traditional Arabic", "editor_font_size": "24"}
    )
    out = tmp_path / "gen.docx"
    r = client.post("/api/docx/export", json={"document_id": did, "target_path": str(out)})
    assert r.status_code == 200, r.text
    assert r.json()["stats"]["paragraphs"] == 4
    d = Document(str(out))
    assert [p.text for p in d.paragraphs] == text.split("\n")  # blank line preserved
    members = read_members(out)
    doc_xml = parse_xml(members["word/document.xml"])
    ps = list(doc_xml.iter(qn("w:p")))
    assert len(ps) == 4 and all(p.find(qn("w:pPr")).find(qn("w:bidi")) is not None for p in ps)
    runs = list(doc_xml.iter(qn("w:r")))
    assert runs and all(r.find(qn("w:rPr")).find(qn("w:rtl")) is not None for r in runs)
    assert doc_xml.find(".//" + qn("w:sectPr")).find(qn("w:bidi")) is not None
    styles = parse_xml(members["word/styles.xml"])
    normal = [s for s in styles.iter(qn("w:style")) if s.get(qn("w:styleId")) == "Normal"][0]
    rpr = normal.find(qn("w:rPr"))
    fonts = rpr.find(qn("w:rFonts"))
    assert fonts.get(qn("w:cs")) == "Traditional Arabic" == fonts.get(qn("w:ascii"))
    assert rpr.find(qn("w:szCs")).get(qn("w:val")) == "36"  # 24px -> 18pt -> 36 half-points
    assert rpr.find(qn("w:lang")).get(qn("w:bidi")) == "ar-SA"
    assert normal.find(qn("w:pPr")).find(qn("w:bidi")) is not None
    assert d.core_properties.title == "عنوان"


def test_generate_docx_diacritizes_raw_content_and_default_font(client, tmp_path):
    did = _doc(client, raw_content="كتاب جديد\nذهب الولد")
    out = tmp_path / "gen2.docx"
    r = client.post(
        "/api/docx/export", json={"document_id": did, "target_path": str(out), "level": "full"}
    )
    assert r.status_code == 200 and r.json()["stats"]["words_diacritized"] == 4
    paras = [p.text for p in Document(str(out)).paragraphs]
    assert paras[0] == "كَتابُ جَديدُ"
    assert [strip_marks(p) for p in paras] == ["كتاب جديد", "ذهب الولد"]
    styles = parse_xml(read_members(out)["word/styles.xml"])
    normal = [s for s in styles.iter(qn("w:style")) if s.get(qn("w:styleId")) == "Normal"][0]
    assert normal.find(qn("w:rPr")).find(qn("w:rFonts")).get(qn("w:cs")) == "Amiri"


def test_generate_docx_errors_and_atomicity(client, tmp_path, monkeypatch):
    did = _doc(client, raw_content="كتاب")
    out = tmp_path / "x.docx"

    def deny(a, b):
        raise PermissionError("locked")

    monkeypatch.setattr(os, "replace", deny)
    r = client.post("/api/docx/export", json={"document_id": did, "target_path": str(out)})
    assert r.status_code == 423 and r.json()["error_code"] == "ERR_DOCX_WRITE_LOCKED"
    assert not out.exists() and not [p for p in tmp_path.iterdir() if p.name.startswith(".~")]
    import errno

    def full(a, b):
        raise OSError(errno.ENOSPC, "full")

    monkeypatch.setattr(os, "replace", full)
    r = client.post("/api/docx/export", json={"document_id": did, "target_path": str(out)})
    assert r.status_code == 507
    assert not [p for p in tmp_path.iterdir() if p.name.startswith(".~")]


def test_imported_document_still_uses_source_file(client, tmp_path):
    src = tmp_path / "rich.docx"
    build_rich_docx(src)
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    did = client.post("/api/docx/import", json={"path": str(src), "project_id": pid}).json()[
        "document_id"
    ]
    out = tmp_path / "o.docx"
    r = client.post("/api/docx/export", json={"document_id": did, "target_path": str(out)})
    assert r.status_code == 200 and "textboxes" in r.json()["stats"]["parts"]
    assert "word/footnotes.xml" in read_members(out)  # original structure kept


# ---------- 7.3.3: marks typed after a tab are preserved ----------
def test_marks_after_tab_are_preserved(client, tmp_path):
    src = build_tab_marks_docx(tmp_path / "tab.docx")
    out = tmp_path / "tab_out.docx"
    r = client.post("/api/docx/export", json={"source_path": str(src), "target_path": str(out)})
    assert r.status_code == 200, r.text
    before = parse_xml(read_members(src)["word/document.xml"])
    after = parse_xml(read_members(out)["word/document.xml"])
    tb = [t.text for t in before.iter(W_T)]
    ta = [t.text for t in after.iter(W_T)]
    assert [strip_marks(x) for x in ta] == [strip_marks(x) for x in tb]
    # the user's shadda / fatha typed right after each tab are still there, at the same place
    assert ta[1].startswith("ّ") and ta[2].startswith("َ")
    assert sum(c in MARKS for c in "".join(ta)) > sum(c in MARKS for c in "".join(tb))
    assert [c.tag.split("}")[1] for c in after.iter(qn("w:r")).__next__()] == [
        "t",
        "tab",
        "t",
        "tab",
        "t",
    ]
