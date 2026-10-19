"""«تصدير Word مع تعقّب التغييرات»: proofreading and diacritization changes as real Word
revisions over the original file. Accept-all == the processed text, reject-all == the original."""

from __future__ import annotations

import zipfile

import pytest
from docx import Document
from lxml import etree

from app.docx.tracked import diff_changes, splice, validate_tracked_docx
from app.engine.normalize import strip_marks

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def Q(tag: str) -> str:
    return f"{{{W}}}{tag}"


ORIG = ["ذهب الولد الي المدرسه", "هذا كتاب جميل", "الفقرة الثالثة هنا"]
# proofreading: الي -> إلى, المدرسه -> المدرسة; then tashkeel on every word
PROC = [
    "ذَهَبَ الوَلَدُ إِلَى المَدْرَسَةِ",
    "هَذَا كِتَابٌ جَمِيلٌ",
    "الفَقْرَةُ الثَّالِثَةُ هُنَا",
]
# the same with proofreading only (no marks)
PROOF_ONLY = ["ذهب الولد إلى المدرسة", "هذا كتاب جميل", "الفقرة الثالثة هنا"]
# tashkeel only (words whose letters were proofread keep their original form)
DIAC_ONLY = [
    "ذَهَبَ الوَلَدُ الي المدرسه",
    "هَذَا كِتَابٌ جَمِيلٌ",
    "الفَقْرَةُ الثَّالِثَةُ هُنَا",
]


def _build(path):
    d = Document()
    p = d.add_paragraph()
    p.add_run("ذهب ").bold = True
    p.add_run("الولد الي ")
    p.add_run("المدرسه").italic = True
    for t in ORIG[1:]:
        d.add_paragraph(t)
    d.save(str(path))
    return path


def _document_xml(path):
    with zipfile.ZipFile(path) as z:
        return etree.fromstring(z.read("word/document.xml"))


def _resolve(root, accept: bool) -> list[str]:
    """Paragraph texts after accepting (or rejecting) every revision."""
    out = []
    for p in root.iter(Q("p")):
        parts = []
        for el in p.iter(Q("t"), Q("delText")):
            anc = {a.tag for a in el.iterancestors()}
            if el.tag == Q("t") and (accept or Q("ins") not in anc):
                parts.append(el.text or "")
            elif el.tag == Q("delText") and not accept:
                parts.append(el.text or "")
        out.append("".join(parts))
    return out


def _import(client, src):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    r = client.post("/api/docx/import", json={"path": str(src), "project_id": pid})
    return r.json()["document_id"]


def _export(client, did, out, **kw):
    r = client.post(
        "/api/docx/export-tracked", json={"document_id": did, "target_path": str(out), **kw}
    )
    assert r.status_code == 200, r.text
    return r.json()["stats"], _document_xml(out)


@pytest.fixture
def doc(client, tmp_path):
    src = _build(tmp_path / "a.docx")
    did = _import(client, src)
    client.put(f"/api/documents/{did}", json={"diacritized_content": "\n".join(PROC)})
    return did, src


def test_both_accept_equals_processed_reject_equals_original(client, tmp_path, doc):
    did, _src = doc
    stats, root = _export(client, did, tmp_path / "o.docx")
    assert _resolve(root, True) == PROC
    assert _resolve(root, False) == ORIG
    assert stats["author"] == "مِزَاج" and stats["skipped_paragraphs"] == 0
    ins, dels = list(root.iter(Q("ins"))), list(root.iter(Q("del")))
    # 2 proofread words (their marks ride in the same revision) + 8 mark-only words
    assert stats["proofread_revisions"] == 2
    assert stats["diacritic_revisions"] == 8
    assert stats["revisions"] == len(ins) == len(dels) == 10
    for e in (*ins, *dels):
        assert e.get(Q("author")) == "مِزَاج" and e.get(Q("date")).endswith("Z")
    ids = [e.get(Q("id")) for e in (*ins, *dels)]
    assert len(ids) == len(set(ids))
    assert validate_tracked_docx(tmp_path / "o.docx") == {"ins": 10, "del": 10}


def test_per_word_not_per_character(client, tmp_path, doc):
    did, _ = doc
    _, root = _export(client, did, tmp_path / "o.docx")
    for d in root.iter(Q("del")):
        assert len("".join(t.text for t in d.iter(Q("delText"))).split()) == 1
    for i in root.iter(Q("ins")):
        assert len("".join(t.text for t in i.iter(Q("t"))).split()) == 1


def test_formatting_survives(client, tmp_path, doc):
    did, _ = doc
    _, root = _export(client, did, tmp_path / "o.docx")
    first = next(root.iter(Q("p")))
    runs = list(first.iter(Q("r")))
    bold = [r for r in runs if r.find(f"{Q('rPr')}/{Q('b')}") is not None]
    ital = [r for r in runs if r.find(f"{Q('rPr')}/{Q('i')}") is not None]
    assert bold and ital
    assert "ذهب" in "".join(t.text for r in bold for t in r.iter(Q("t"), Q("delText")))
    assert Document(str(tmp_path / "o.docx")).paragraphs  # still opens in python-docx


def test_proofread_only(client, tmp_path, doc):
    did, _ = doc
    stats, root = _export(client, did, tmp_path / "o.docx", scope="proofread")
    assert _resolve(root, True) == PROOF_ONLY
    assert _resolve(root, False) == ORIG
    assert stats["revisions"] == 2 and stats["diacritic_revisions"] == 0
    assert len(list(root.iter(Q("ins")))) == 2


def test_diacritics_only(client, tmp_path, doc):
    did, _ = doc
    stats, root = _export(client, did, tmp_path / "o.docx", scope="diacritics")
    assert _resolve(root, True) == DIAC_ONLY
    assert _resolve(root, False) == ORIG
    assert stats["proofread_revisions"] == 0 and stats["diacritic_revisions"] == 8


def test_char_granularity_revises_clusters(client, tmp_path, doc):
    did, _ = doc
    stats, root = _export(client, did, tmp_path / "o.docx", granularity="char")
    assert _resolve(root, True) == PROC and _resolve(root, False) == ORIG
    assert stats["diacritic_revisions"] > 8  # per letter cluster instead of per word


def test_unchanged_document_has_no_revisions(client, tmp_path):
    src = _build(tmp_path / "a.docx")
    did = _import(client, src)
    stats, root = _export(client, did, tmp_path / "o.docx")
    assert stats["revisions"] == 0 and not list(root.iter(Q("ins")))
    assert _resolve(root, False) == ORIG


def test_text_document_without_source_uses_original_text(client, tmp_path):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    did = client.post(
        "/api/documents",
        json={"project_id": pid, "title": "T", "raw_content": "\n".join(PROC)},
    ).json()["id"]
    out = tmp_path / "o.docx"
    r = client.post("/api/docx/export-tracked", json={"document_id": did, "target_path": str(out)})
    assert r.status_code == 422  # original_text is required without a source file
    stats, root = _export(client, did, out, original_text="\n".join(ORIG))
    assert _resolve(root, True) == PROC and _resolve(root, False) == ORIG
    assert stats["revisions"] == 10


def test_existing_revision_paragraph_is_written_untracked(tmp_path):
    """A paragraph that already holds revisions cannot nest ours: counted, never corrupted."""
    from app.docx.tracked import process_docx_edits_tracked

    d = Document()
    d.add_paragraph("هذا كتاب جميل")
    d.save(str(tmp_path / "a.docx"))
    with zipfile.ZipFile(tmp_path / "a.docx") as z:
        parts = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(parts["word/document.xml"])
    r = next(root.iter(Q("r")))
    ins = etree.Element(
        Q("ins"), {Q("id"): "5", Q("author"): "X", Q("date"): "2020-01-01T00:00:00Z"}
    )
    r.addprevious(ins)
    ins.append(r)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    with zipfile.ZipFile(tmp_path / "b.docx", "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    stats = process_docx_edits_tracked(tmp_path / "b.docx", tmp_path / "o.docx", "هَذَا كِتَابٌ جَمِيلٌ")
    assert stats["untracked_paragraphs"] == 1 and stats["revisions"] == 0
    assert strip_marks(_resolve(_document_xml(tmp_path / "o.docx"), True)[0]) == "هذا كتاب جميل"


def test_diff_changes_unit():
    old, new = "ذهب الي المدرسه، ثم عاد", "ذَهَبَ إلى المدرسةِ ثم عاد."
    for scope in ("both", "proofread", "diacritics"):
        ch = diff_changes(old, new, scope)
        assert all(0 <= c.a <= c.b <= len(old) for c in ch)
        assert all(a.b <= b.a for a, b in zip(ch, ch[1:], strict=False))
    assert splice(old, diff_changes(old, new, "both")) == new
    assert strip_marks(splice(old, diff_changes(old, new, "proofread"))) == strip_marks(new)
    only = splice(old, diff_changes(old, new, "diacritics"))
    assert only.startswith("ذَهَبَ الي") and only.endswith("ثم عاد")
    assert {c.kind for c in diff_changes(old, new, "both")} == {"proof", "diac"}
    assert splice("أ ب", diff_changes("أ ب", "أ، ب")) == "أ، ب"  # pure insertion
    assert splice("أ، ب", diff_changes("أ، ب", "أ ب")) == "أ ب"  # pure deletion


def test_malformed_output_is_rejected(tmp_path):
    from app.core.errors import MizajError

    d = Document()
    d.add_paragraph("نص")
    d.save(str(tmp_path / "a.docx"))
    with zipfile.ZipFile(tmp_path / "a.docx") as z:
        parts = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(parts["word/document.xml"])
    run = next(root.iter(Q("t"))).getparent()
    dele = etree.Element(
        Q("del"), {Q("id"): "9", Q("author"): "x", Q("date"): "2020-01-01T00:00:00Z"}
    )
    run.addprevious(dele)
    dele.append(run)  # w:t inside w:del: wrong element
    parts["word/document.xml"] = etree.tostring(root)
    with zipfile.ZipFile(tmp_path / "bad.docx", "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    with pytest.raises(MizajError):
        validate_tracked_docx(tmp_path / "bad.docx")


def test_text_document_proofread_scope_and_punctuation(client, tmp_path):
    old = "محكمة ابو حمص ، فى ان المدعي اقام الدعوى ."
    new = "مَحْكَمَةُ أَبُو حَمْصٍ، فِي أَنَّ الْمُدَّعِي أَقَامَ الدَّعْوَى."
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    did = client.post(
        "/api/documents", json={"project_id": pid, "title": "T", "raw_content": new}
    ).json()["id"]
    stats, root = _export(client, did, tmp_path / "o.docx", scope="proofread", original_text=old)
    assert stats["proofread_revisions"] == 6 and stats["skipped_paragraphs"] == 0
    assert _resolve(root, False) == [old]
    assert strip_marks(_resolve(root, True)[0]) == strip_marks(new)
    assert len(list(root.iter(Q("ins")))) == 4  # the two pure deletions have no w:ins
