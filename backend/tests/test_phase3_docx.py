import errno
import hashlib
import os
import time
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx_factory import (
    CODE_TEXT,
    COMMENT_TEXT,
    ENDNOTE_TEXT,
    FOOTER_TEXTS,
    FOOTNOTE_TEXT,
    HEADER_TEXTS,
    INSTR_TEXT,
    SDT_TEXT,
    TEXTBOX_TEXT,
    build_rich_docx,
    read_members,
)
from lxml import etree

from app.core.database import connect
from app.docx.extractor import discover_parts, iter_units
from app.docx.mixed import is_code_run, tokenize, verify_mixed
from app.docx.signature import run_signature, segments
from app.engine import config
from app.engine.mock_engine import MockEngine
from app.engine.normalize import MARKS, strip_marks

FIX = Path(__file__).parent / "fixtures"
W_T = qn("w:t")
ENVELOPE_KEYS = {"error_code", "message_ar", "message_en", "details"}
MODIFIABLE = ("word/document.xml", "word/header", "word/footer", "word/footnotes", "word/endnotes")
mock = MockEngine()


def vz(text):
    return mock.vowelize_text(text)


@pytest.fixture
def rich(tmp_path):
    return build_rich_docx(tmp_path / "rich.docx")


def export(client, src, dst, **kw):
    return client.post(
        "/api/docx/export", json={"source_path": str(src), "target_path": str(dst), **kw}
    )


def xml_parts(members):
    return {
        n: parse_xml(b)
        for n, b in members.items()
        if n.startswith("word/") and n.endswith(".xml") and b"<w:" in b
    }


def texts(root):
    return [t.text or "" for t in root.iter(W_T)]


def neutral(root):
    clone = etree.fromstring(etree.tostring(root))
    for t in clone.iter(W_T):
        t.text = None
    return etree.tostring(clone, method="c14n")


def para_texts(root):
    return ["".join(t.text or "" for t in p.iter(W_T)) for p in root.iter(qn("w:p"))]


def has_marks(s):
    return any(c in MARKS for c in s)


# ---------- unit: signature / mixed / extractor ----------
def test_signature_is_the_seven_tuple_and_segments(rich):
    doc = Document(str(rich))
    p0 = doc.paragraphs[0]
    sigs = [run_signature(r) for r in p0.runs]
    assert all(len(s) == 7 for s in sigs)
    assert sigs[0] == sigs[1]  # "الك" / "تاب"
    assert sigs[3] != sigs[2]  # bold run differs
    segs = segments(p0._p)
    assert [len(s.runs) for s in segs] == [3, 1, 1]  # [الك+تاب+" جميل "], [مدر bold], [سة]
    # field / tab / footnote-reference runs are never grouped
    for p in doc.paragraphs:
        for seg in segments(p._p):
            for r in seg.runs:
                assert all(c.tag in (W_T, qn("w:rPr")) for c in r)


def test_mixed_tokenize_and_code_run():
    toks = tokenize("اشتريت iPhone 16 الجديد!")
    assert "".join(t for _, t in toks) == "اشتريت iPhone 16 الجديد!"
    assert [a for a, _ in toks] == [True, False, True, False]
    assert verify_mixed("كتاب 12", "كَتابُ 12") and not verify_mixed("كتاب 12", "كَتابُ 13")
    r = parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:rPr><w:rFonts w:cs=" fira CODE "/></w:rPr><w:t>x</w:t></w:r>'
    )
    assert is_code_run(r)


def test_extractor_parts_deduped_and_categories(rich):
    doc = Document(str(rich))
    parts = discover_parts(doc, include_comments=False)
    names = [n for n, _ in parts]
    assert len(names) == len(set(names))
    kinds = [k for _, k in parts]
    assert kinds.count("headers") == 3 and kinds.count("footers") == 3
    assert "footnotes" in kinds and "endnotes" in kinds and "comments" not in kinds
    assert ("word/comments.xml", "comments") in discover_parts(doc, True)
    root = doc.element
    cats = [u.category for u in iter_units(root, "body")]
    assert cats.count("textboxes") == 2  # Choice + Fallback
    assert cats.count("sdt") == 1


# ---------- full export on the rich document ----------
@pytest.fixture
def rich_out(client, rich, tmp_path):
    out = tmp_path / "out" / "rich_out.docx"
    out.parent.mkdir()
    before = hashlib.sha256(rich.read_bytes()).hexdigest()
    r = export(client, rich, out)
    assert r.status_code == 200, r.text
    assert hashlib.sha256(rich.read_bytes()).hexdigest() == before  # source untouched
    return {"src": read_members(rich), "dst": read_members(out), "resp": r.json(), "path": out}


def test_stripped_text_identical_for_every_wt_node(rich_out):
    src, dst = xml_parts(rich_out["src"]), xml_parts(rich_out["dst"])
    assert src.keys() == dst.keys()
    for name in src:
        a, b = texts(src[name]), texts(dst[name])
        assert len(a) == len(b), name
        assert [strip_marks(x) for x in b] == a, name


def test_all_non_text_xml_identical(rich_out):
    """Every rPr/pPr/rsid/attribute/non-text node is untouched: trees match once w:t text is
    blanked out."""
    src, dst = xml_parts(rich_out["src"]), xml_parts(rich_out["dst"])
    for name in src:
        assert neutral(src[name]) == neutral(dst[name]), name


def test_unrelated_zip_members_byte_identical(rich_out):
    s, d = rich_out["src"], rich_out["dst"]
    assert list(s) == list(d)  # same members, same order
    modified = {n for n in s if s[n] != d[n]}
    assert modified and all(n.startswith(MODIFIABLE) for n in modified)
    assert "word/comments.xml" not in modified
    for n in (
        "[Content_Types].xml",
        "word/_rels/document.xml.rels",
        "word/styles.xml",
        "word/settings.xml",
        "docProps/core.xml",
        "word/comments.xml",
    ):
        assert s[n] == d[n], n


def test_split_words_redistributed_by_letter_ownership(rich_out):
    root = parse_xml(rich_out["dst"]["word/document.xml"])
    p1 = list(root.iter(qn("w:p")))[0]
    ts = texts(p1)
    # "الك"+"تاب" -> mock word 'اَلكتابُ'; every letter stays in its original w:t
    assert ts[0] == "اَلك" and ts[1] == "تابُ"
    assert "".join(ts) == vz("الكتاب جميل مدرسة")
    # word split across DIFFERENT formatting (bold "مدر" / plain "سة") keeps run formatting
    assert ts[3] == "مَدر" and ts[4] == "سةُ"
    runs = list(p1.iter(qn("w:r")))
    assert runs[3].find(qn("w:rPr")).find(qn("w:b")) is not None
    assert runs[4].find(qn("w:rPr")).find(qn("w:b")) is None


def test_special_runs_untouched_and_structures(rich_out):
    d = rich_out["dst"]
    root = parse_xml(d["word/document.xml"])
    instr = [e.text for e in root.iter(qn("w:instrText"))]
    assert instr == [INSTR_TEXT]
    all_t = texts(root)
    assert CODE_TEXT in all_t and "اختبار" in all_t  # code-font runs skipped
    assert any(t.startswith("اشتريت") for t in all_t) is False  # mixed paragraph got marks
    mixed = [p for p in para_texts(root) if "iPhone" in p][0]
    assert "iPhone 16 Pro" in mixed and "1200 USD" in mixed and "https://example.com/a" in mixed
    assert verify_mixed(
        "اشتريت جهاز iPhone 16 Pro الجديد بمبلغ 1200 USD من https://example.com/a", mixed
    )
    # tab / br survive inside the run; text around is diacritized
    tabbed = [r for r in root.iter(qn("w:r")) if r.find(qn("w:tab")) is not None][0]
    assert [c.tag.split("}")[1] for c in tabbed] == ["t", "tab", "t", "br", "t"]
    tabbed_t = [t.text for t in tabbed.iter(W_T)]
    assert tabbed_t[:2] == [vz("قبل"), vz("بعد")] and has_marks(tabbed_t[2])
    # hyperlink, bookmarks, fldSimple, footnote reference kept
    assert len(list(root.iter(qn("w:hyperlink")))) == 1
    assert len(list(root.iter(qn("w:bookmarkStart")))) == 1
    assert len(list(root.iter(qn("w:fldSimple")))) == 1
    assert len(list(root.iter(qn("w:footnoteReference")))) == 1
    link = list(root.iter(qn("w:hyperlink")))[0]
    assert "".join(texts(link)) == vz("رابط الموقع")


def test_tables_headers_footers_notes_textbox_sdt(rich_out):
    d = rich_out["dst"]
    doc_root = parse_xml(d["word/document.xml"])
    src_root = parse_xml(rich_out["src"]["word/document.xml"])
    assert len(list(doc_root.iter(qn("w:tbl")))) == len(list(src_root.iter(qn("w:tbl")))) == 2
    cells = [
        vz(x) for x in ("رأس مدمج للجدول", "عمود ثالث", "خلية عادية", "جدول متداخل", "خلية داخلية")
    ]
    body_texts = para_texts(doc_root)
    for c in cells:
        assert c in body_texts
    assert len(list(doc_root.iter(qn("w:gridSpan")))) == 1  # merged cell kept
    hdr = {n: para_texts(parse_xml(b)) for n, b in d.items() if n.startswith("word/header")}
    ftr = {n: para_texts(parse_xml(b)) for n, b in d.items() if n.startswith("word/footer")}
    assert sorted(x for v in hdr.values() for x in v if x) == sorted(
        vz(t) for t in HEADER_TEXTS.values()
    )
    assert sorted(x for v in ftr.values() for x in v if x) == sorted(
        vz(t) for t in FOOTER_TEXTS.values()
    )
    assert vz(FOOTNOTE_TEXT) in [x.strip() for x in para_texts(parse_xml(d["word/footnotes.xml"]))]
    assert vz(ENDNOTE_TEXT) in [x.strip() for x in para_texts(parse_xml(d["word/endnotes.xml"]))]
    tb = [e for e in doc_root.iter(qn("w:txbxContent"))]
    assert len(tb) == 2
    assert [vz(TEXTBOX_TEXT)] * 2 == ["".join(texts(e)) for e in tb]  # Choice == Fallback
    assert vz(SDT_TEXT) in body_texts
    parts = rich_out["resp"]["stats"]["parts"]
    assert parts["textboxes"] == 2 and parts["sdt"] == 1 and parts["headers"] == 3
    assert parts["footers"] == 3 and parts["endnotes"] == 1 and parts["footnotes"] == 1
    assert parts["body"] > 8 and "comments" not in parts
    assert rich_out["resp"]["stats"]["skipped_code_runs"] == 2


def test_comments_untouched_by_default_and_opt_in(client, rich, tmp_path):
    out = tmp_path / "c.docx"
    assert export(client, rich, out).status_code == 200
    assert read_members(out)["word/comments.xml"] == read_members(rich)["word/comments.xml"]
    out2 = tmp_path / "c2.docx"
    r = export(client, rich, out2, include_comments=True)
    assert r.status_code == 200 and r.json()["stats"]["parts"]["comments"] == 1
    c = parse_xml(read_members(out2)["word/comments.xml"])
    assert texts(c) == [vz(COMMENT_TEXT)]


def test_output_opens_with_python_docx(rich_out):
    doc = Document(str(rich_out["path"]))
    assert len(doc.tables) == 1 and len(doc.paragraphs) >= 9
    assert doc.sections[0].different_first_page_header_footer
    assert has_marks(doc.paragraphs[0].text)
    src_doc = Document(str(rich_out["path"].parent.parent / "rich.docx"))
    assert [len(p.runs) for p in doc.paragraphs] == [len(p.runs) for p in src_doc.paragraphs]


def test_levels(client, rich, tmp_path):
    outs = {}
    for lvl in ("full", "irab_only", "minimal"):
        o = tmp_path / f"{lvl}.docx"
        assert export(client, rich, o, level=lvl).status_code == 200
        outs[lvl] = parse_xml(read_members(o)["word/document.xml"])
    base = parse_xml(read_members(rich)["word/document.xml"])
    for lvl, root in outs.items():
        assert [strip_marks(t) for t in texts(root)] == texts(base), lvl
    n = {k: sum(c in MARKS for t in texts(v) for c in t) for k, v in outs.items()}
    assert n["full"] > n["irab_only"] > 0 and n["full"] > n["minimal"] > 0
    first = para_texts(outs["irab_only"])[0]
    assert first == "الكتابُ جميلُ مدرسةُ"  # only last-letter marks
    assert para_texts(outs["minimal"])[0] == "الكتاب جميل مدرسة"  # no 3-letter word -> bare
    assert vz("جدول") in "".join(para_texts(outs["full"]))
    # minimal (mock rule): 3-letter words only, e.g. "عمود ثالث" has none; "رأس" is 3 letters
    assert "رَأسُ" in "".join(para_texts(outs["minimal"]))


# ---------- the real fixtures ----------
def test_fixture_sample_ar(client, tmp_path):
    src = FIX / "sample_ar.docx"
    h = hashlib.sha256(src.read_bytes()).hexdigest()
    out = tmp_path / "ar.docx"
    r = export(client, src, out)
    assert r.status_code == 200
    assert hashlib.sha256(src.read_bytes()).hexdigest() == h
    s, d = read_members(src), read_members(out)
    assert list(s) == list(d)
    assert {n for n in s if s[n] != d[n]} <= {
        "word/document.xml",
        "word/header1.xml",
        "word/footer1.xml",
    }
    doc = Document(str(out))
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 2
    body = parse_xml(d["word/document.xml"])
    p = para_texts(body)[0]
    assert p == vz("العلم نور يضيء طريق المستقبل في كل زمان ومكان.")
    assert "كَتابُ" == vz("كتاب")
    assert vz("مدرسة") in para_texts(body)
    for name in xml_parts(s):
        assert neutral(xml_parts(s)[name]) == neutral(xml_parts(d)[name])
    assert r.json()["stats"]["parts"]["body"] == 5


def test_fixture_sample_mixed(client, tmp_path):
    out = tmp_path / "mixed.docx"
    r = export(client, FIX / "sample_mixed.docx", out)
    assert r.status_code == 200 and r.json()["stats"]["skipped_code_runs"] == 1
    root = parse_xml(read_members(out)["word/document.xml"])
    src = parse_xml(read_members(FIX / "sample_mixed.docx")["word/document.xml"])
    t = texts(root)
    assert t[1] == "calculate_total(price=100)" == texts(src)[1]
    assert "iPhone 16 Pro" in t[0] and "1200 USD" in t[0] and "https://example.com/store" in t[0]
    assert has_marks(t[0]) and has_marks(t[2])
    assert neutral(root) == neutral(src)


@pytest.mark.skipif(not config.model_present(), reason="CATT model not downloaded")
def test_real_accurate_export_sample_mixed(client, monkeypatch, tmp_path):
    monkeypatch.delenv("MIZAJ_ENGINE_MODE", raising=False)
    out = tmp_path / "real.docx"
    r = export(client, FIX / "sample_mixed.docx", out, mode="accurate")
    assert r.status_code == 200, r.text
    root = parse_xml(read_members(out)["word/document.xml"])
    src = parse_xml(read_members(FIX / "sample_mixed.docx")["word/document.xml"])
    t = texts(root)
    assert [strip_marks(x) for x in t] == texts(src)
    assert "iPhone 16 Pro" in t[0] and "1200 USD" in t[0] and t[1] == texts(src)[1]
    assert sum(c in MARKS for c in t[0]) > 20
    assert neutral(root) == neutral(src)


# ---------- atomic save / errors ----------
def leftovers(d):
    return [p.name for p in d.iterdir() if p.name.startswith(".~mizaj_tmp")]


def test_locked_target_returns_423_and_cleans_temp(client, rich, tmp_path, monkeypatch):
    out = tmp_path / "locked.docx"

    def deny(a, b):
        raise PermissionError("in use")

    monkeypatch.setattr(os, "replace", deny)
    r = export(client, rich, out)
    assert r.status_code == 423 and r.json()["error_code"] == "ERR_DOCX_WRITE_LOCKED"
    assert set(r.json()) == ENVELOPE_KEYS
    assert not out.exists() and leftovers(tmp_path) == []


def test_disk_full_returns_507(client, rich, tmp_path, monkeypatch):
    out = tmp_path / "full.docx"

    def nospace(a, b):
        raise OSError(errno.ENOSPC, "No space left on device")

    monkeypatch.setattr(os, "replace", nospace)
    r = export(client, rich, out)
    assert r.status_code == 507 and r.json()["error_code"] == "ERR_DISK_FULL"
    assert not out.exists() and leftovers(tmp_path) == []


def test_existing_target_replaced_atomically_and_temp_name(client, rich, tmp_path, monkeypatch):
    out = tmp_path / "exist.docx"
    out.write_bytes(b"old")
    seen = []
    real = os.replace
    monkeypatch.setattr(os, "replace", lambda a, b: (seen.append(Path(a).name), real(a, b))[1])
    assert export(client, rich, out).status_code == 200
    assert seen and seen[0].startswith(".~mizaj_tmp.") and seen[0].endswith(".docx")
    assert zipfile.is_zipfile(out) and leftovers(tmp_path) == []


def test_concurrent_exports_use_distinct_temp_files(client, rich, tmp_path):
    assert export(client, rich, tmp_path / "warm.docx").status_code == 200  # engine loaded
    seen = set()
    real = os.replace
    from concurrent.futures import ThreadPoolExecutor

    import app.docx.service as svc

    orig = svc.os.replace
    svc.os.replace = lambda a, b: (seen.add(Path(a).name), real(a, b))[1]
    try:
        with ThreadPoolExecutor(4) as ex:
            res = list(
                ex.map(
                    lambda i: export(client, rich, tmp_path / f"c{i}.docx").status_code, range(4)
                )
            )
    finally:
        svc.os.replace = orig
    assert res == [200] * 4 and len(seen) == 4


def test_corrupt_inputs_return_422(client, tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"this is not a zip")
    r = export(client, bad, tmp_path / "o.docx")
    assert r.status_code == 422 and r.json()["error_code"] == "ERR_DOCX_CORRUPT"
    assert set(r.json()) == ENVELOPE_KEYS
    notdocx = tmp_path / "nodoc.docx"
    with zipfile.ZipFile(notdocx, "w") as z:
        z.writestr("hello.txt", "hi")
    assert export(client, notdocx, tmp_path / "o.docx").json()["error_code"] == "ERR_DOCX_CORRUPT"
    broken = tmp_path / "broken.docx"
    members = read_members(FIX / "sample_ar.docx")
    members["word/document.xml"] = b"<w:document><oops>"
    with zipfile.ZipFile(broken, "w") as z:
        for n, b in members.items():
            z.writestr(n, b)
    assert export(client, broken, tmp_path / "o.docx").json()["error_code"] == "ERR_DOCX_CORRUPT"
    trunc = tmp_path / "trunc.docx"
    trunc.write_bytes((FIX / "sample_ar.docx").read_bytes()[:500])
    assert export(client, trunc, tmp_path / "o.docx").status_code == 422
    assert not (tmp_path / "o.docx").exists() and leftovers(tmp_path) == []


# ---------- API: import / export ----------
def test_import_creates_document_and_recent_file(client):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    src = FIX / "sample_mixed.docx"
    r = client.post("/api/docx/import", json={"path": str(src), "project_id": pid})
    assert r.status_code == 201, r.text
    j = r.json()
    assert j["title"] == "sample_mixed" and j["word_count"] > 10
    assert {"paragraphs", "words", "tables"} <= set(j["stats"])
    doc = client.get(f"/api/documents/{j['document_id']}").json()
    assert doc["source_docx_path"] == str(src)
    assert (
        "iPhone 16 Pro" in doc["raw_content"] and "calculate_total(price=100)" in doc["raw_content"]
    )
    assert doc["word_count"] == j["word_count"]
    conn = connect(client.data_dir / "mizaj.db")
    rows = conn.execute("SELECT file_path, document_id FROM recent_files").fetchall()
    assert [(x[0], x[1]) for x in rows] == [(str(src), j["document_id"])]
    r2 = client.post("/api/docx/import", json={"path": str(src), "project_id": pid, "title": "T2"})
    rows = conn.execute("SELECT file_path, document_id FROM recent_files").fetchall()
    assert len(rows) == 1 and rows[0][1] == r2.json()["document_id"]
    conn.close()


def test_import_table_document_text(client):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    r = client.post(
        "/api/docx/import", json={"path": str(FIX / "sample_ar.docx"), "project_id": pid}
    ).json()
    text = client.get(f"/api/documents/{r['document_id']}").json()["raw_content"]
    assert text.splitlines()[0].startswith("العلم نور يضيء طريق المستقبل")
    assert "مدرسة" in text.splitlines() and r["stats"]["tables"] == 1


def test_export_by_document_id(client, tmp_path):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    did = client.post(
        "/api/docx/import", json={"path": str(FIX / "sample_ar.docx"), "project_id": pid}
    ).json()["document_id"]
    out = tmp_path / "byid.docx"
    r = client.post(
        "/api/docx/export",
        json={"document_id": did, "target_path": str(out), "mode": "fast", "level": "full"},
    )
    assert r.status_code == 200
    assert r.json()["target_path"] == str(out) and zipfile.is_zipfile(out)
    assert set(r.json()["stats"]) == {
        "paragraphs",
        "words_diacritized",
        "skipped_code_runs",
        "parts",
    }
    assert set(r.json()["stats"]["parts"]) == {
        "body",
        "headers",
        "footers",
        "footnotes",
        "endnotes",
        "textboxes",
        "sdt",
    }
    assert r.json()["stats"]["words_diacritized"] > 5


def test_api_validation(client, tmp_path, rich):
    ok = str(tmp_path / "x.docx")
    cases = [
        {"source_path": "rel/a.docx", "target_path": ok},
        {"source_path": str(rich), "target_path": str(tmp_path / "x.txt")},
        {"source_path": str(rich), "target_path": "x.docx"},
        {"target_path": ok},
        {"source_path": str(rich), "document_id": 1, "target_path": ok},
        {"source_path": str(rich), "target_path": ok, "level": "bogus"},
        {"source_path": str(rich), "target_path": str(tmp_path / "nodir" / "x.docx")},
        {"source_path": str(rich), "target_path": str(rich)},
    ]
    for body in cases:
        r = client.post("/api/docx/export", json=body)
        assert r.status_code == 422, body
        assert r.json()["error_code"] == "ERR_VALIDATION"
    r = client.post(
        "/api/docx/export", json={"source_path": str(tmp_path / "no.docx"), "target_path": ok}
    )
    assert r.status_code == 404 and r.json()["error_code"] == "ERR_NOT_FOUND"
    assert (
        client.post("/api/docx/export", json={"document_id": 99, "target_path": ok}).status_code
        == 404
    )
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    assert (
        client.post("/api/docx/import", json={"path": "a.docx", "project_id": pid}).status_code
        == 422
    )
    assert (
        client.post(
            "/api/docx/import", json={"path": str(tmp_path / "n.docx"), "project_id": pid}
        ).status_code
        == 404
    )
    assert (
        client.post("/api/docx/import", json={"path": str(rich), "project_id": 999}).status_code
        == 404
    )


def test_export_does_not_block_health(client, rich, tmp_path, monkeypatch):
    orig = MockEngine.diacritize

    def slow(self, words, lookup=None):
        time.sleep(0.05)
        return orig(self, words, lookup)

    monkeypatch.setattr(MockEngine, "diacritize", slow)
    import asyncio

    import httpx

    from app.main import create_app

    app = create_app(client.data_dir / "other")

    async def scenario():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://127.0.0.1"
        ) as c:
            task = asyncio.create_task(
                c.post(
                    "/api/docx/export",
                    json={"source_path": str(rich), "target_path": str(tmp_path / "h.docx")},
                    timeout=60,
                )
            )
            await asyncio.sleep(0.2)
            assert not task.done()
            t0 = time.perf_counter()
            assert (await c.get("/health")).status_code == 200
            lat = time.perf_counter() - t0
            return lat, await task

    lat, resp = asyncio.run(scenario())
    assert lat < 0.5 and resp.status_code == 200
